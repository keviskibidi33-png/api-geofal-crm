from __future__ import annotations
 
import io
import json # Added by user
import logging
import os
import re
import zipfile
from copy import copy
from datetime import date, datetime # Added datetime by user
from pathlib import Path # Corrected from 'from pathlib import os'
from typing import Any
 
import asyncio
import requests # Moved by user
from fastapi import FastAPI, HTTPException, Header, Response, Depends, Request # Added Response, Depends; kept Header
from fastapi.exceptions import RequestValidationError
from fastapi.exception_handlers import request_validation_exception_handler
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse, JSONResponse # Added JSONResponse
from dotenv import load_dotenv
from openpyxl import load_workbook
from openpyxl.worksheet.pagebreak import Break
from openpyxl.utils.cell import get_column_letter, range_boundaries
from pydantic import BaseModel, Field
import psycopg2
import psycopg2.errors # Explicitly import errors module
from psycopg2.extras import RealDictCursor, Json

# Importar el nuevo exportador XML
# from app.xlsx_direct_v2 import export_xlsx_direct
# from app.programacion_export import export_programacion_xlsx # Removed
from app.modules.recepcion.router import router as recepciones_router
from app.modules.cotizacion.router import router as cotizacion_router
from app.modules.programacion.router import router as programacion_router
from app.modules.verificacion.router import router as verificacion_router
from app.modules.compresion.router import router as compresion_router
from app.modules.tracing.router import router as tracing_router
from app.modules.humedad.router import router as humedad_router
from app.modules.cont_humedad.router import router as cont_humedad_router
from app.modules.planas.router import router as planas_router
from app.modules.caras.router import router as caras_router
from app.modules.cbr.router import router as cbr_router
from app.modules.proctor.router import router as proctor_router
from app.modules.llp.router import router as llp_router
from app.modules.gran_suelo.router import router as gran_suelo_router
from app.modules.gran_agregado.router import router as gran_agregado_router
from app.modules.abra.router import router as abra_router
from app.modules.abrass.router import router as abrass_router
from app.modules.peso_unitario.router import router as peso_unitario_router
from app.modules.tamiz.router import router as tamiz_router
from app.modules.equi_arena.router import router as equi_arena_router
from app.modules.ge_fino.router import router as ge_fino_router
from app.modules.ge_grueso.router import router as ge_grueso_router
from app.modules.cd.router import router as cd_router
from app.modules.ph.router import router as ph_router
from app.modules.cloro_soluble.router import router as cloro_soluble_router
from app.modules.sales_solubles.router import router as sales_solubles_router
from app.modules.sulfatos_solubles.router import router as sulfatos_solubles_router
from app.modules.compresion_no_confinada.router import router as compresion_no_confinada_router
from app.modules.cont_mat_organica.router import router as cont_mat_organica_router
from app.modules.terrones_fino_grueso.router import router as terrones_fino_grueso_router
from app.modules.azul_metileno.router import router as azul_metileno_router
from app.modules.part_livianas.router import router as part_livianas_router
from app.modules.imp_organicas.router import router as imp_organicas_router
from app.modules.sul_magnesio.router import router as sul_magnesio_router
from app.modules.angularidad.router import router as angularidad_router
from app.modules.ingenieria_archivos.router import router as ingenieria_archivos_router
from app.modules.correlativos.router import router as correlativos_router
from app.modules.control_informes.router import router as control_informes_router
from app.modules.seguimiento_cliente_comercial.router import router as seguimiento_comercial_router
from app.modules.publicidad_geofal.router import router as publicidad_geofal_router
from app.modules.control_probetas.router import router as control_probetas_router
from app.modules.densidad_huantar.router import router as densidad_huantar_router
from app.modules.recepcion.models import Base as RecepcionBase
from app.modules.verificacion.models import Base as VerificacionBase
from app.modules.tracing.models import Trazabilidad
from app.modules.compresion.models import EnsayoCompresion, ItemCompresion
from app.modules.humedad.models import HumedadEnsayo
from app.modules.cont_humedad.models import ContHumedadEnsayo
from app.modules.planas.models import PlanasEnsayo
from app.modules.caras.models import CarasEnsayo
from app.modules.cbr.models import CBREnsayo
from app.modules.proctor.models import ProctorEnsayo
from app.modules.llp.models import LLPEnsayo
from app.modules.gran_suelo.models import GranSueloEnsayo
from app.modules.gran_agregado.models import GranAgregadoEnsayo
from app.modules.abra.models import AbraEnsayo
from app.modules.abrass.models import AbrassEnsayo
from app.modules.peso_unitario.models import PesoUnitarioEnsayo
from app.modules.tamiz.models import TamizEnsayo
from app.modules.equi_arena.models import EquiArenaEnsayo
from app.modules.ge_fino.models import GeFinoEnsayo
from app.modules.ge_grueso.models import GeGruesoEnsayo
from app.modules.cd.models import CDEnsayo
from app.modules.ph.models import PHEnsayo
from app.modules.cloro_soluble.models import CloroSolubleEnsayo
from app.modules.sales_solubles.models import SalesSolublesEnsayo
from app.modules.sulfatos_solubles.models import SulfatosSolublesEnsayo
from app.modules.compresion_no_confinada.models import CompresionNoConfinadaEnsayo
from app.modules.cont_mat_organica.models import ContMatOrganicaEnsayo
from app.modules.terrones_fino_grueso.models import TerronesFinoGruesoEnsayo
from app.modules.azul_metileno.models import AzulMetilenoEnsayo
from app.modules.densidad_huantar.models import DensidadHuantarEnsayo
from app.modules.part_livianas.models import PartLivianasEnsayo
from app.modules.imp_organicas.models import ImpOrganicasEnsayo
from app.modules.sul_magnesio.models import SulMagnesioEnsayo
from app.modules.angularidad.models import AngularidadEnsayo
from app.modules.ingenieria_archivos.models import IngenieriaArchivo
from app.modules.correlativos.models import CorrelativoReserva, CorrelativoTurno
from app.modules.control_informes.models import (
    ControlEnsayoCatalogo,
    ControlEnsayoCounter,
    ControlInforme,
    ControlInformeDetalle,
)
from app.modules.seguimiento_cliente_comercial.models import SeguimientoClienteComercial
from app.modules.publicidad_geofal.models import PublicidadGeofal
from app.database import engine
from app.auth import JWTAuthMiddleware
from app.utils.http_client import http_get, http_patch

logger = logging.getLogger(__name__)

# Ensure tables are created safely
try:
    RecepcionBase.metadata.create_all(bind=engine)
    VerificacionBase.metadata.create_all(bind=engine)
    # Compression tables will be created via migration or explicitly:
    from app.database import Base as MainBase
    MainBase.metadata.create_all(bind=engine)
    
    # Programmatic migrations to ensure database schema alignment in all environments
    try:
        from sqlalchemy import text
        with engine.begin() as conn:
            # Migration 044: Add control_probetas permissions to admin roles
            conn.execute(text("""
                UPDATE role_definitions
                SET permissions = jsonb_set(permissions, '{control_probetas}', '{"read": true, "write": true, "delete": true}'::jsonb, true)
                WHERE role_id IN ('admin', 'admin_general');
            """))
            # Update for Oficina Técnica roles to enable Control Probetas
            conn.execute(text("""
                UPDATE role_definitions
                SET permissions = jsonb_set(permissions, '{control_probetas}', '{"read": true, "write": true, "delete": true}'::jsonb, true)
                WHERE role_id IN ('oficina_tecnica', 'oficina_tecnica_humedad', 'oficina_tecnica_humedad_tipificador', 'oficina_tecnica_sup');
            """))
            logger.info("Programmatic migration 044 applied successfully (or was already applied).")
    except Exception as perm_err:
        logger.warning("Could not apply migration 044 permissions: %s", perm_err)

    try:
        from sqlalchemy import text
        with engine.begin() as conn:
            # Migration 045: Add control probetas columns to muestras_concreto table
            conn.execute(text("ALTER TABLE public.muestras_concreto ADD COLUMN IF NOT EXISTS elemento VARCHAR(50) DEFAULT '-';"))
            conn.execute(text("ALTER TABLE public.muestras_concreto ADD COLUMN IF NOT EXISTS fosa VARCHAR(20) DEFAULT '-';"))
            conn.execute(text("ALTER TABLE public.muestras_concreto ADD COLUMN IF NOT EXISTS densidad VARCHAR(50) DEFAULT '-';"))
            conn.execute(text("ALTER TABLE public.muestras_concreto ADD COLUMN IF NOT EXISTS status_ensayo VARCHAR(50) DEFAULT '-';"))
            conn.execute(text("ALTER TABLE public.muestras_concreto ADD COLUMN IF NOT EXISTS status_entrega VARCHAR(50) DEFAULT '-';"))
            conn.execute(text("ALTER TABLE public.muestras_concreto ADD COLUMN IF NOT EXISTS fecha_entrega VARCHAR(50) DEFAULT '-';"))
            conn.execute(text("ALTER TABLE public.muestras_concreto ADD COLUMN IF NOT EXISTS es_control_probetas BOOLEAN DEFAULT FALSE;"))
            conn.execute(text("CREATE INDEX IF NOT EXISTS idx_muestras_concreto_es_control_probetas ON public.muestras_concreto (es_control_probetas);"))
            conn.execute(text("CREATE INDEX IF NOT EXISTS idx_muestras_concreto_fecha_rotura ON public.muestras_concreto (fecha_rotura);"))
            conn.execute(text("CREATE INDEX IF NOT EXISTS idx_muestras_concreto_recepcion_id ON public.muestras_concreto (recepcion_id);"))
            conn.execute(text("NOTIFY pgrst, 'reload schema';"))
            logger.info("Programmatic migration 045 applied successfully (or was already applied).")
    except Exception as col_err:
        logger.warning("Could not apply migration 045 columns: %s", col_err)

    try:
        from sqlalchemy import text
        with engine.begin() as conn:
            # Migration 046: Add densidad_huantar permissions to admin and technical roles
            conn.execute(text("""
                UPDATE role_definitions
                SET permissions = jsonb_set(permissions, '{densidad_huantar}', '{"read": true, "write": true, "delete": true}'::jsonb, true)
                WHERE role_id IN ('admin', 'admin_general', 'oficina_tecnica', 'oficina_tecnica_humedad', 'oficina_tecnica_humedad_tipificador', 'oficina_tecnica_sup', 'jefe_laboratorio', 'tecnico', 'tecnico_suelos');
            """))
            logger.info("Programmatic migration 046 applied successfully.")
    except Exception as perm_err:
        logger.warning("Could not apply migration 046 permissions: %s", perm_err)

except Exception as e:
    logger.warning("Could not create database tables on startup (DB might be offline): %s", e)

# --- Pydantic Models for Roles & Permissions ---

class ModulePermission(BaseModel):
    read: bool = False
    write: bool = False
    delete: bool = False

class RolePermissions(BaseModel):
    clientes: ModulePermission | None = None
    proyectos: ModulePermission | None = None
    cotizadora: ModulePermission | None = None
    programacion: ModulePermission | None = None
    recepcion: ModulePermission | None = None
    verificacion_muestras: ModulePermission | None = None
    compresion: ModulePermission | None = None
    tracing: ModulePermission | None = None
    control_probetas: ModulePermission | None = None
    densidad_huantar: ModulePermission | None = None
    humedad: ModulePermission | None = None
    cont_humedad: ModulePermission | None = None
    planas: ModulePermission | None = None
    caras: ModulePermission | None = None
    cbr: ModulePermission | None = None
    proctor: ModulePermission | None = None
    llp: ModulePermission | None = None
    gran_suelo: ModulePermission | None = None
    gran_agregado: ModulePermission | None = None
    abra: ModulePermission | None = None
    abrass: ModulePermission | None = None
    peso_unitario: ModulePermission | None = None
    tamiz: ModulePermission | None = None
    equi_arena: ModulePermission | None = None
    ge_fino: ModulePermission | None = None
    ge_grueso: ModulePermission | None = None
    cd: ModulePermission | None = None
    ph: ModulePermission | None = None
    cloro_soluble: ModulePermission | None = None
    sales_solubles: ModulePermission | None = None
    sulfatos_solubles: ModulePermission | None = None
    compresion_no_confinada: ModulePermission | None = None
    cont_mat_organica: ModulePermission | None = None
    terrones_fino_grueso: ModulePermission | None = None
    azul_metileno: ModulePermission | None = None
    part_livianas: ModulePermission | None = None
    imp_organicas: ModulePermission | None = None
    sul_magnesio: ModulePermission | None = None
    angularidad: ModulePermission | None = None
    ingenieria_archivos: ModulePermission | None = None
    control_informes: ModulePermission | None = None
    correlativos: ModulePermission | None = None
    usuarios: ModulePermission | None = None
    auditoria: ModulePermission | None = None
    configuracion: ModulePermission | None = None
    laboratorio: ModulePermission | None = None
    oficina_tecnica: ModulePermission | None = None
    comercial: ModulePermission | None = None
    administracion: ModulePermission | None = None
    permisos: ModulePermission | None = None

class RoleDefinition(BaseModel):
    role_id: str
    label: str
    description: str | None = None
    permissions: RolePermissions
    is_system: bool = False
    created_at: datetime | None = None
    updated_at: datetime | None = None

class RoleUpdate(BaseModel):
    label: str | None = None
    description: str | None = None
    permissions: RolePermissions | None = None

class UserPermissionOverrideUpdate(BaseModel):
    enabled: bool = True
    permissions: dict[str, ModulePermission] = Field(default_factory=dict)

class HeartbeatRequest(BaseModel):
    user_id: str


class DashboardSearchItem(BaseModel):
    id: str
    type: str
    title: str
    subtitle: str


class DashboardSearchResponse(BaseModel):
    data: list[DashboardSearchItem]


class DashboardNotification(BaseModel):
    id: str
    type: str
    severity: str = "warning"
    title: str
    message: str
    status: str = "open"
    created_at: datetime | None = None
    acknowledged_at: datetime | None = None
    metadata: dict[str, Any] = Field(default_factory=dict)

app = FastAPI(title="quotes-service")

@app.exception_handler(Exception)
async def global_exception_handler(request: Request, exc: Exception):
    """
    Manejador global de excepciones para asegurar que incluso los errores no manejados
    devuelvan una respuesta JSON válida y no rompan los encabezados CORS.
    """
    logger.error("Unhandled exception: %s", exc, exc_info=True)
    return JSONResponse(
        status_code=500,
        content={
            "detail": "Internal Server Error",
            "message": str(exc) if os.getenv("ALLOW_INSECURE_DEV_AUTH") == "true" else "Error interno del servidor",
            "path": request.url.path
        }
    )

load_dotenv(Path(__file__).resolve().parents[1] / ".env", override=False)

if (os.getenv("QUOTES_DATABASE_URL") or "").strip() == "":
    os.environ.pop("QUOTES_DATABASE_URL", None)


def _db_disabled() -> bool:
    return (os.getenv("QUOTES_DISABLE_DB") or "").strip().lower() in {"1", "true", "yes", "on"}


def _get_cors_origins() -> list[str]:
    origins = [
        "http://localhost:8474",
        "http://127.0.0.1:8474",
        "http://localhost:3000", 
        "http://localhost:3001", 
        "http://localhost:3002",
        "http://localhost:3003", # Compresión CRM
        "http://localhost:3004", 
        "http://localhost:3005", # Verificación CRM
        "http://localhost:3006",
        "http://localhost:3007",
        "http://localhost:3009", # Proctor CRM (Vite local)
        "http://localhost:3010", # LLP CRM (Vite local)
        "http://localhost:3011", # Gran Suelo CRM (Vite local)
        "http://localhost:3012", # Gran Agregado CRM (Vite local)
        "http://localhost:3013", # EquiArena CRM (Vite local)
        "http://localhost:3014", # GE Fino CRM (Vite local)
        "http://localhost:3015", # GE Grueso CRM (Vite local)
        "http://localhost:3016", # ABRA CRM (Vite local)
        "http://localhost:3017", # Peso Unitario CRM (Vite local)
        "http://localhost:3018", # Tamiz CRM (Vite local)
        "http://localhost:3019", # ABRASS CRM (Vite local)
        "http://localhost:3020", # Contenido Humedad CRM (Vite local)
        "http://localhost:3021", # Planas CRM (Vite local)
        "http://localhost:3022", # Caras CRM (Vite local)
        "http://localhost:3023", # CD CRM (Vite local)
        "http://localhost:3024", # PH CRM (Vite local)
        "http://localhost:3025", # Cloro Soluble CRM (Vite local)
        "http://localhost:3026", # Sales Solubles CRM (Vite local)
        "http://localhost:3027", # Sulfatos Solubles CRM (Vite local)
        "http://localhost:3028", # Compresion No Confinada CRM (Vite local)
        "http://localhost:3029", # Ensayos Especiales CRM (Vite local)
        "http://localhost:5173", # Cotizador
        "http://localhost:5174",
        "http://localhost:5175", # Compresion (Vite)
        "http://localhost:5176", # LLP CRM (Vite local alternate)
        "http://127.0.0.1:3000", 
        "http://127.0.0.1:3009",
        "http://127.0.0.1:3010",
        "http://127.0.0.1:3011",
        "http://127.0.0.1:3012",
        "http://127.0.0.1:3013",
        "http://127.0.0.1:3014",
        "http://127.0.0.1:3015",
        "http://127.0.0.1:3016",
        "http://127.0.0.1:3017",
        "http://127.0.0.1:3018",
        "http://127.0.0.1:3019",
        "http://127.0.0.1:3020",
        "http://127.0.0.1:3030",
        "http://127.0.0.1:3021",
        "http://127.0.0.1:3022",
        "http://127.0.0.1:3023",
        "http://127.0.0.1:3024",
        "http://127.0.0.1:3025",
        "http://127.0.0.1:3026",
        "http://127.0.0.1:3027",
        "http://127.0.0.1:3028",
        "http://127.0.0.1:3029",
        "http://127.0.0.1:3030",
        "http://127.0.0.1:5173",
        "http://127.0.0.1:5174",
        "http://127.0.0.1:5175",
        "http://127.0.0.1:5176",
        "https://crm.geofal.com.pe",
        "https://recepcion.geofal.com.pe",
        "https://cotizador.geofal.com.pe",
        "https://programacion.geofal.com.pe",
        "https://compresion.geofal.com.pe",
        "https://laboratorio.geofal.com.pe", # Added just in case
        "https://humedad.geofal.com.pe",
        "https://cbr.geofal.com.pe",
        "https://proctor.geofal.com.pe",
        "https://llp.geofal.com.pe",
        "https://gran-suelo.geofal.com.pe",
        "https://gran-agregado.geofal.com.pe",
        "https://equiarena.geofal.com.pe",
        "https://equi-arena.geofal.com.pe",
        "https://ge-fino.geofal.com.pe",
        "https://ge-grueso.geofal.com.pe",
        "https://abra.geofal.com.pe",
        "https://abrass.geofal.com.pe",
        "https://peso-unitario.geofal.com.pe",
        "https://tamiz.geofal.com.pe",
        "https://contenido-humedad.geofal.com.pe",
        "https://planas.geofal.com.pe",
        "https://caras.geofal.com.pe",
        "https://cd.geofal.com.pe",
        "https://ph.geofal.com.pe",
        "https://cloro-soluble.geofal.com.pe",
        "https://sales-solubles.geofal.com.pe",
        "https://sulfatos-solubles.geofal.com.pe",
        "https://compresion-no-confinada.geofal.com.pe",
        "https://ensayos-especiales.geofal.com.pe",
        "https://comp.noconfinada.geofal.com.pe",
        "https://verificacion.geofal.com.pe",
        "https://verifiacion.geofal.com.pe",
    ]
    raw = os.getenv("QUOTES_CORS_ORIGINS")
    if raw:
        if raw == "*":
            return ["*"]
        extra = [o.strip() for o in raw.split(",") if o.strip()]
        origins.extend(extra)
    return list(set(origins))

 
 
# Determine CORS origins
_origins = _get_cors_origins()
# If origins are set specifically, allow credentials. If it's "*", we cannot.
_allow_creds = "*" not in _origins and len(_origins) > 0

# Regex to allow any geofal subdomain (prod) and local dev environments
_origin_regex = r"https?://([a-zA-Z0-9-]+\.)*geofal\.com\.pe(:\d+)?|https?://localhost(:\d+)?|https?://127\.0\.0\.1(:\d+)?"

logger.info("CORS configured: origins=%s allow_credentials=%s", _origins, _allow_creds)

# Force HTTPS behind reverse proxies using X-Forwarded-Proto header
@app.middleware("http")
async def force_https_behind_proxy(request: Request, call_next):
    if request.headers.get("x-forwarded-proto") == "https":
        request.scope["scheme"] = "https"
    response = await call_next(request)
    return response

# JWT Auth Middleware
app.add_middleware(JWTAuthMiddleware)

# CORS Middleware must be outermost so even 401/500 responses include CORS headers
app.add_middleware(
    CORSMiddleware,
    allow_origins=_origins,
    allow_origin_regex=_origin_regex,
    allow_credentials=_allow_creds,
    allow_methods=["GET", "POST", "PUT", "DELETE", "OPTIONS", "PATCH"],
    allow_headers=["*"],
    expose_headers=[
        "Content-Disposition",
        "X-Humedad-Id",
        "X-Cont-Humedad-Id",
        "X-Planas-Id",
        "X-Caras-Id",
        "X-CBR-Id",
        "X-Proctor-Id",
        "X-LLP-Id",
        "X-Gran-Suelo-Id",
        "X-Gran-Agregado-Id",
        "X-ABRA-Id",
        "X-ABRASS-Id",
        "X-Peso-Unitario-Id",
        "X-Tamiz-Id",
        "X-Equi-Arena-Id",
        "X-Ge-Fino-Id",
        "X-Ge-Grueso-Id",
        "X-CD-Id",
        "X-PH-Id",
        "X-CL-Id",
        "X-SS-Id",
        "X-SULF-Id",
        "X-CNC-Id",
        "X-CMO-Id",
        "X-TFG-Id",
        "X-AZM-Id",
        "X-PLV-Id",
        "X-IMP-Id",
        "X-SMAG-Id",
        "X-ANG-Id",
        "X-Storage-Object-Key",
        "X-Control-Probetas-Id",
    ],
    max_age=3600,
)
 
 
def _get_database_url() -> str:
    url = os.getenv("QUOTES_DATABASE_URL")
    if not url:
        return f"postgresql://{os.getenv('DB_USER', 'directus')}:{os.getenv('DB_PASSWORD', 'directus')}@{os.getenv('DB_HOST', 'postgres')}:{os.getenv('DB_PORT', '5432')}/{os.getenv('DB_DATABASE', 'directus')}"
    return url

def _has_database_url() -> bool:
    if _db_disabled(): return False
    url = (os.getenv("QUOTES_DATABASE_URL") or "").strip()
    if url: return True
    return bool(os.getenv("DB_HOST"))

def _get_connection():
    dsn = _get_database_url()
    return psycopg2.connect(dsn, connect_timeout=3)

@app.exception_handler(ValueError)
async def value_error_handler(request, exc):
    """Catch-all for simple validation logic errors in service layer"""
    return JSONResponse(
        status_code=400,
        content={"detail": str(exc)},
    )


@app.exception_handler(RequestValidationError)
async def request_validation_error_handler(request: Request, exc: RequestValidationError):
    """Manejador de validación de Pydantic para devolver respuestas limpias en español."""
    errors = exc.errors()
    logger.warning("Validation error on %s %s: %s", request.method, request.url.path, errors)
    
    friendly_messages = []
    for err in errors:
        loc = err.get("loc", [])
        msg = err.get("msg", "")
        err_type = err.get("type", "")
        
        field_name = str(loc[-1]) if loc else "campo"
        
        # Traducir los tipos de errores comunes de Pydantic
        if err_type == "string_too_long":
            max_len = err.get("ctx", {}).get("max_length", 50)
            msg_es = f"no debe superar los {max_len} caracteres"
        elif err_type == "missing":
            msg_es = "es un campo obligatorio"
        elif err_type == "value_error":
            msg_es = msg.replace("Value error,", "").strip()
        else:
            msg_es = msg
            
        # Identificar si el error pertenece a una muestra específica
        if len(loc) >= 4 and loc[1] == "muestras_verificadas":
            muestra_idx = loc[2]
            friendly_field = "Código LEM" if field_name == "codigo_lem" else ("Tipo de Testigo" if field_name == "tipo_testigo" else field_name)
            friendly_messages.append(f"Muestra #{muestra_idx + 1}: El '{friendly_field}' {msg_es}.")
        else:
            friendly_messages.append(f"El campo '{field_name}' {msg_es}.")
            
    return JSONResponse(
        status_code=422,
        content={
            "error": "ValidationError",
            "message": " | ".join(friendly_messages),
            "code": 422,
            "detail": errors
        }
    )


 
@app.get("/")
@app.get("/health")
async def health_check():
    return {"status": "ok", "service": "quotes-service", "db": _has_database_url()}


@app.get("/debug-db")
async def debug_db():
    """Verify database connection and schema for troubleshooting"""
    if not _has_database_url():
        return {"error": "DATABASE_URL not set"}
    try:
        conn = _get_connection()
        with conn.cursor() as cur:
            cur.execute("SELECT version();")
            version = cur.fetchone()[0]
            
            cur.execute("SELECT table_name FROM information_schema.tables WHERE table_schema='public'")
            tables = [r[0] for r in cur.fetchall()]
            
            return {
                "status": "connected",
                "version": version,
                "tables": tables,
                "dsn_start": _get_database_url().split('@')[-1] if '@' in _get_database_url() else "local"
            }
    except Exception as e:
        import traceback
        tb = traceback.format_exc()
        print(f"DEBUG-DB Error: {e}\n{tb}")
        return {"status": "error", "message": str(e), "traceback": tb}
    finally:
        if 'conn' in locals() and conn:
            conn.close()



# Include Routers
app.include_router(recepciones_router)
app.include_router(cotizacion_router)
app.include_router(programacion_router)
app.include_router(verificacion_router)
app.include_router(compresion_router)
app.include_router(tracing_router)
app.include_router(humedad_router)
app.include_router(cont_humedad_router)
app.include_router(planas_router)
app.include_router(caras_router)
app.include_router(cbr_router)
app.include_router(proctor_router)
app.include_router(llp_router)
app.include_router(gran_suelo_router)
app.include_router(gran_agregado_router)
app.include_router(abra_router)
app.include_router(abrass_router)
app.include_router(peso_unitario_router)
app.include_router(tamiz_router)
app.include_router(equi_arena_router)
app.include_router(ge_fino_router)
app.include_router(ge_grueso_router)
app.include_router(cd_router)
app.include_router(ph_router)
app.include_router(cloro_soluble_router)
app.include_router(sales_solubles_router)
app.include_router(sulfatos_solubles_router)
app.include_router(compresion_no_confinada_router)
app.include_router(cont_mat_organica_router)
app.include_router(terrones_fino_grueso_router)
app.include_router(azul_metileno_router)
app.include_router(part_livianas_router)
app.include_router(imp_organicas_router)
app.include_router(sul_magnesio_router)
app.include_router(angularidad_router)
app.include_router(ingenieria_archivos_router)
app.include_router(correlativos_router)
app.include_router(control_informes_router)
app.include_router(seguimiento_comercial_router)
app.include_router(publicidad_geofal_router)
app.include_router(control_probetas_router)
app.include_router(densidad_huantar_router)

# Note: All legacy endpoints for Quotes and Programacion have been moved to their respective modules.
# Check app/modules/cotizacion and app/modules/programacion.
 






 
 
# ===================== USER PROFILE ENDPOINT =====================

@app.get("/user/me")
async def get_current_user(authorization: str = Header(None)):
    """Get current user profile from Directus token"""
    if not authorization:
        return {"data": None}
    
    try:
        # Forward request to Directus
        directus_url = os.getenv("DIRECTUS_URL", "http://directus:8055")
        resp = await asyncio.get_event_loop().run_in_executor(
            None,
            lambda: __import__('urllib.request', fromlist=['urlopen']).urlopen(
                __import__('urllib.request', fromlist=['Request']).Request(
                    f"{directus_url}/users/me",
                    headers={"Authorization": authorization}
                )
            )
        )
        import json as json_module
        data = json_module.loads(resp.read().decode())
        user = data.get('data', {})
        return {
            "data": {
                "id": user.get('id'),
                "first_name": user.get('first_name'),
                "last_name": user.get('last_name'),
                "email": user.get('email'),
                "phone": user.get('phone') or user.get('telefono'),
            }
        }
    except Exception as e:
        print(f"Error fetching user: {e}")
        return {"data": None}


# ===================== CLIENTS & PROJECTS ENDPOINTS =====================

@app.get("/clientes")
async def get_clientes(search: str = ""):
    """Get clients list with optional search - reads from CRM clientes table"""
    if not _has_database_url():
        return {"data": []}
    
    try:
        conn = _get_connection()
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            if search:
                cur.execute("""
                    SELECT DISTINCT ON (c.id)
                        c.id, c.empresa, c.ruc, c.direccion,
                        COALESCE(con.nombre, c.nombre) as col_contacto,
                        COALESCE(con.email, c.email) as col_email,
                        COALESCE(con.telefono, c.telefono) as col_telefono
                    FROM clientes c
                    LEFT JOIN contactos con ON c.id = con.cliente_id
                    WHERE (c.nombre ILIKE %s OR c.empresa ILIKE %s OR c.email ILIKE %s OR con.nombre ILIKE %s)
                    AND c.deleted_at IS NULL
                    ORDER BY c.id, con.es_principal DESC
                    LIMIT 20
                """, (f"%{search}%", f"%{search}%", f"%{search}%", f"%{search}%"))
            else:
                cur.execute("""
                    SELECT id, nombre as col_contacto, email as col_email, telefono as col_telefono, 
                           empresa, ruc, direccion
                    FROM clientes 
                    WHERE deleted_at IS NULL
                    ORDER BY nombre LIMIT 50
                """)
            results = cur.fetchall()
            # Map to cotizador expected format (B2B professional format)
            # empresa = company name (primary), nombre = contact person (secondary)
            mapped = [{
                'id': str(r['id']),
                'nombre': r.get('empresa') or r.get('col_contacto', ''),  # empresa as main client name
                'contacto': r.get('col_contacto', ''),  # matched contact or main contact
                'email': r.get('col_email', ''),
                'telefono': r.get('col_telefono', ''),
                'ruc': r.get('ruc', ''),
                'direccion': r.get('direccion', ''),
            } for r in results]
            return {"data": mapped}
    except Exception as e:
        import traceback
        print(f"Error in get_clientes: {e}")
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=str(e))
    finally:
        if 'conn' in locals() and conn:
            conn.close()


@app.post("/clientes")
async def create_cliente(data: dict):
    """Create a new client - uses same columns as CRM"""
    if not _has_database_url():
        raise HTTPException(status_code=400, detail="Database not configured")
    
    try:
        conn = _get_connection()
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            cur.execute("""
                INSERT INTO clientes (nombre, email, telefono, empresa, ruc, estado, sector, direccion)
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s)
                RETURNING id, nombre, email, telefono, empresa, ruc, estado, sector, direccion
            """, (
                data.get('contacto', '') or data.get('nombre', ''),
                data.get('email', ''),
                data.get('telefono', ''),
                data.get('nombre', ''),
                data.get('ruc', ''),
                'prospecto',
                'General',
                data.get('direccion', '')
            ))
            result = cur.fetchone()
            conn.commit()
            # Map back to cotizador format
            mapped = {
                'id': str(result['id']),
                'nombre': result.get('empresa', ''),
                'contacto': result.get('nombre', ''),
                'email': result.get('email', ''),
                'telefono': result.get('telefono', ''),
                'ruc': result.get('ruc', ''),
                'direccion': result.get('direccion', ''),
            }
            return {"data": mapped}
    except psycopg2.errors.UniqueViolation as e:
        if 'conn' in locals() and conn:
            conn.rollback()
        # Friendly error message
        if 'clientes_ruc_key' in str(e) or 'ruc' in str(e):
             raise HTTPException(status_code=409, detail="Ya existe un cliente registrado con este número de RUC. Por favor verifique.")
        raise HTTPException(status_code=409, detail="Error de duplicidad: Ya existe un registro con estos datos.")
    except Exception as e:
        import traceback
        print(f"Error in create_cliente: {e}")
        traceback.print_exc()
        if 'conn' in locals() and conn:
            conn.rollback()
        raise HTTPException(status_code=500, detail=str(e))
    finally:
        if 'conn' in locals() and conn:
            conn.close()


@app.get("/proyectos")
async def get_proyectos(cliente_id: str = None, search: str = ""):
    """Get projects list, optionally filtered by client - reads from CRM proyectos table"""
    if not _has_database_url():
        return {"data": []}
    
    try:
        conn = _get_connection()
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            query = """
                SELECT 
                    p.id, p.nombre, p.descripcion, p.cliente_id, p.created_at, p.direccion, p.ubicacion, 
                    c.empresa as cliente_nombre,
                    v.full_name as vendedor_nombre, v.phone as vendedor_telefono
                FROM proyectos p
                LEFT JOIN clientes c ON p.cliente_id = c.id
                LEFT JOIN perfiles v ON p.vendedor_id = v.id
                WHERE p.deleted_at IS NULL
            """
            params = []

            if cliente_id:
                query += " AND p.cliente_id = %s"
                params.append(cliente_id)
            
            if search:
                query += " AND p.nombre ILIKE %s"
                params.append(f"%{search}%")
            
            query += " ORDER BY p.nombre LIMIT 50"
            
            cur.execute(query, tuple(params))
            results = cur.fetchall()
            
            # Map results to ensure JSON serializability (handle datetime and UUID)
            mapped = []
            for r in results:
                mapped.append({
                    'id': str(r['id']),
                    'nombre': r['nombre'],
                    'direccion': r.get('direccion', ''),
                    'ubicacion': r.get('ubicacion', ''),
                    'descripcion': r.get('descripcion', ''),
                    'cliente_id': str(r['cliente_id']),
                    'cliente_nombre': r.get('cliente_nombre', ''),
                    'vendedor_nombre': r.get('vendedor_nombre', ''),
                    'vendedor_telefono': r.get('vendedor_telefono', ''),
                    'created_at': r['created_at'].isoformat() if r.get('created_at') else None
                })
            
            return {"data": mapped}
    except Exception as e:
        import traceback
        print(f"Error in get_proyectos: {e}")
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=str(e))
    finally:
        if 'conn' in locals() and conn:
            conn.close()


@app.get("/dashboard/search", response_model=DashboardSearchResponse)
async def dashboard_search(q: str = "", limit: int = 10):
    """Búsqueda rápida unificada para header del CRM."""
    if not _has_database_url():
        return DashboardSearchResponse(data=[])

    safe_limit = max(1, min(limit, 20))
    query_text = (q or "").strip()

    try:
        conn = _get_connection()
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            if not query_text:
                cur.execute(
                    """
                    WITH recent_clients AS (
                        SELECT
                            c.id::text AS id,
                            'cliente'::text AS type,
                            COALESCE(NULLIF(c.empresa, ''), NULLIF(c.nombre, ''), 'Sin nombre') AS title,
                            CASE
                                WHEN COALESCE(c.ruc, '') <> '' THEN 'RUC: ' || c.ruc
                                ELSE 'Cliente reciente'
                            END AS subtitle,
                            1 AS section_order,
                            c.created_at AS sort_at
                        FROM clientes c
                        WHERE c.deleted_at IS NULL
                        ORDER BY c.created_at DESC NULLS LAST
                        LIMIT 3
                    ),
                    recent_projects AS (
                        SELECT
                            p.id::text AS id,
                            'proyecto'::text AS type,
                            COALESCE(NULLIF(p.nombre, ''), 'Sin nombre') AS title,
                            'Estado: ' || COALESCE(NULLIF(p.estado, ''), 'N/A') AS subtitle,
                            2 AS section_order,
                            p.created_at AS sort_at
                        FROM proyectos p
                        WHERE p.deleted_at IS NULL
                        ORDER BY p.created_at DESC NULLS LAST
                        LIMIT 2
                    ),
                    recent_quotes AS (
                        SELECT
                            c.id::text AS id,
                            'cotizacion'::text AS type,
                            COALESCE(NULLIF(c.numero, ''), 'Sin número') AS title,
                            CASE
                                WHEN c.total IS NOT NULL THEN 'S/ ' || trim(to_char(c.total, 'FM9999999990.00'))
                                ELSE 'Cotización reciente'
                            END AS subtitle,
                            3 AS section_order,
                            c.created_at AS sort_at
                        FROM cotizaciones c
                        ORDER BY c.created_at DESC NULLS LAST
                        LIMIT 2
                    )
                    SELECT id, type, title, subtitle
                    FROM (
                        SELECT * FROM recent_clients
                        UNION ALL
                        SELECT * FROM recent_projects
                        UNION ALL
                        SELECT * FROM recent_quotes
                    ) search_results
                    ORDER BY section_order, sort_at DESC NULLS LAST, title
                    LIMIT %s
                    """,
                    (safe_limit,),
                )
            else:
                like_query = f"%{query_text}%"
                is_numeric_query = query_text.isdigit()
                looks_like_quote = "COT" in query_text.upper() or is_numeric_query

                cur.execute(
                    """
                    WITH client_matches AS (
                        SELECT DISTINCT ON (c.id)
                            c.id::text AS id,
                            'cliente'::text AS type,
                            COALESCE(NULLIF(c.empresa, ''), NULLIF(c.nombre, ''), 'Sin nombre') AS title,
                            CASE
                                WHEN COALESCE(c.ruc, '') <> '' THEN 'RUC: ' || c.ruc
                                WHEN COALESCE(con.email, c.email, '') <> '' THEN COALESCE(con.email, c.email)
                                ELSE 'Sin contacto'
                            END AS subtitle,
                            1 AS section_order,
                            c.created_at AS sort_at
                        FROM clientes c
                        LEFT JOIN contactos con ON con.cliente_id = c.id
                        WHERE c.deleted_at IS NULL
                          AND (
                              c.nombre ILIKE %s
                              OR c.empresa ILIKE %s
                              OR c.email ILIKE %s
                              OR c.ruc ILIKE %s
                              OR con.nombre ILIKE %s
                          )
                        ORDER BY c.id, con.es_principal DESC NULLS LAST, c.created_at DESC NULLS LAST
                        LIMIT 5
                    ),
                    project_matches AS (
                        SELECT
                            p.id::text AS id,
                            'proyecto'::text AS type,
                            COALESCE(NULLIF(p.nombre, ''), 'Sin nombre') AS title,
                            'Estado: ' || COALESCE(NULLIF(p.estado, ''), 'N/A') AS subtitle,
                            2 AS section_order,
                            p.created_at AS sort_at
                        FROM proyectos p
                        WHERE p.deleted_at IS NULL
                          AND p.nombre ILIKE %s
                        ORDER BY p.created_at DESC NULLS LAST, p.nombre
                        LIMIT 5
                    ),
                    quote_matches AS (
                        SELECT
                            c.id::text AS id,
                            'cotizacion'::text AS type,
                            COALESCE(NULLIF(c.numero, ''), 'Sin número') AS title,
                            CASE
                                WHEN c.total IS NOT NULL THEN 'S/ ' || trim(to_char(c.total, 'FM9999999990.00'))
                                WHEN COALESCE(c.cliente_nombre, '') <> '' THEN c.cliente_nombre
                                ELSE 'Sin monto'
                            END AS subtitle,
                            3 AS section_order,
                            c.created_at AS sort_at
                        FROM cotizaciones c
                        WHERE %s
                          AND (
                              c.numero ILIKE %s
                              OR c.cliente_nombre ILIKE %s
                          )
                        ORDER BY c.created_at DESC NULLS LAST, c.numero
                        LIMIT 5
                    )
                    SELECT id, type, title, subtitle
                    FROM (
                        SELECT * FROM client_matches
                        UNION ALL
                        SELECT * FROM project_matches
                        UNION ALL
                        SELECT * FROM quote_matches
                    ) search_results
                    ORDER BY section_order, sort_at DESC NULLS LAST, title
                    LIMIT %s
                    """,
                    (
                        like_query,
                        like_query,
                        like_query,
                        like_query,
                        like_query,
                        like_query,
                        looks_like_quote,
                        like_query,
                        like_query,
                        safe_limit,
                    ),
                )

            rows = cur.fetchall()
            items = [
                DashboardSearchItem(
                    id=str(row["id"]),
                    type=str(row["type"]),
                    title=str(row["title"] or "Sin nombre"),
                    subtitle=str(row["subtitle"] or ""),
                )
                for row in rows
            ]

        return DashboardSearchResponse(data=items)
    except Exception as exc:
        logger.exception("Error en dashboard_search q=%s limit=%s", query_text, safe_limit)
        raise HTTPException(status_code=500, detail=f"Error buscando datos del dashboard: {exc}")
    finally:
        if "conn" in locals() and conn:
            conn.close()


@app.post("/proyectos")
async def create_proyecto(data: dict):
    """Create a new project (requires cliente_id)"""
    if not _has_database_url():
        raise HTTPException(status_code=400, detail="Database not configured")
    
    if not data.get('cliente_id'):
        raise HTTPException(status_code=400, detail="cliente_id is required")
    
    try:
        conn = _get_connection()
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            # Extract vendedor_id from data or fallback to None
            vendedor_id = data.get('vendedor_id') or data.get('user_id')
            
            cur.execute("""
                INSERT INTO proyectos (nombre, ubicacion, descripcion, cliente_id, vendedor_id)
                VALUES (%s, %s, %s, %s, %s)
                RETURNING id, nombre, ubicacion, descripcion, cliente_id, vendedor_id
            """, (
                data.get('nombre', ''),
                data.get('ubicacion', ''),
                data.get('descripcion', ''),
                data.get('cliente_id'),
                vendedor_id
            ))
            result = cur.fetchone()
            conn.commit()
            # Convert results to ensure JSON serializability
            mapped = {k: (str(v) if k in ('id', 'cliente_id', 'vendedor_id') else v) for k, v in result.items()}
            return {"data": mapped}
    except Exception as e:
        import traceback
        print(f"Error in create_proyecto: {e}")
        traceback.print_exc()
        if 'conn' in locals() and conn:
            conn.rollback()
        raise HTTPException(status_code=500, detail=str(e))
    finally:
        if 'conn' in locals() and conn:
            conn.close()


# ============================================================================
# CONDICIONES ESPECÍFICAS ENDPOINTS
# ============================================================================

@app.get("/condiciones")
async def get_condiciones(search: str = ""):
    """Get all active specific conditions, optionally filtered by search"""
    if not _has_database_url():
        return {"data": []}
    
    try:
        conn = _get_connection()
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            query = """
                SELECT id, texto, categoria, orden, created_by, created_at
                FROM condiciones_especificas
                WHERE activo = true
            """
            params = []
            
            if search:
                query += " AND texto ILIKE %s"
                params.append(f"%{search}%")
            
            query += " ORDER BY orden ASC, created_at ASC"
            
            cur.execute(query, params)
            results = cur.fetchall()
            # Ensure JSON serializable
            return {"data": [dict(r) for r in results]}
    except Exception as e:
        print(f"Error in get_condiciones: {e}")
        import traceback
        traceback.print_exc()
        return {"data": []}
    finally:
        if 'conn' in locals() and conn:
            conn.close()


@app.post("/condiciones")
async def create_condicion(data: dict):
    """Create a new specific condition"""
    if not _has_database_url():
        raise HTTPException(status_code=400, detail="Database not configured")
    
    try:
        conn = _get_connection()
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            cur.execute("""
                INSERT INTO condiciones_especificas (texto, categoria, orden, created_by, activo)
                VALUES (%s, %s, %s, %s, %s)
                RETURNING id, texto, categoria, orden, created_by, created_at
            """, (
                data.get('texto', ''),
                data.get('categoria', ''),
                data.get('orden', 0),
                data.get('vendedor_id'),  # created_by = vendedor_id
                True
            ))
            result = cur.fetchone()
            conn.commit()
            return {"data": dict(result)}
    except Exception as e:
        import traceback
        print(f"Error in create_condicion: {e}")
        traceback.print_exc()
        if 'conn' in locals() and conn:
            conn.rollback()
        raise HTTPException(status_code=500, detail=str(e))
    finally:
        if 'conn' in locals() and conn:
            conn.close()


@app.put("/condiciones/{condicion_id}")
async def update_condicion(condicion_id: str, data: dict):
    """Update an existing condition"""
    if not _has_database_url():
        raise HTTPException(status_code=400, detail="Database not configured")
    
    try:
        conn = _get_connection()
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            cur.execute("""
                UPDATE condiciones_especificas
                SET texto = %s, categoria = %s, orden = %s, updated_at = NOW()
                WHERE id = %s
                RETURNING id, texto, categoria, orden, created_at, updated_at
            """, (
                data.get('texto', ''),
                data.get('categoria', ''),
                data.get('orden', 0),
                condicion_id
            ))
            result = cur.fetchone()
            if not result:
                raise HTTPException(status_code=404, detail="Condición no encontrada")
            conn.commit()
            return {"data": dict(result)}
    except HTTPException:
        raise
    except Exception as e:
        import traceback
        print(f"Error in update_condicion: {e}")
        traceback.print_exc()
        if 'conn' in locals() and conn:
            conn.rollback()
        raise HTTPException(status_code=500, detail=str(e))
    finally:
        if 'conn' in locals() and conn:
            conn.close()


@app.delete("/condiciones/{condicion_id}")
async def delete_condicion(condicion_id: str):
    """Soft delete a condition (set activo = false)"""
    if not _has_database_url():
        raise HTTPException(status_code=400, detail="Database not configured")
    
    try:
        conn = _get_connection()
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            cur.execute("""
                UPDATE condiciones_especificas
                SET activo = false, updated_at = NOW()
                WHERE id = %s
                RETURNING id
            """, (condicion_id,))
            result = cur.fetchone()
            if not result:
                raise HTTPException(status_code=404, detail="Condición no encontrada")
            conn.commit()
            return {"message": "Condición eliminada", "id": str(result['id'])}
    except HTTPException:
        raise
    except Exception as e:
        import traceback
        print(f"Error in delete_condicion: {e}")
        traceback.print_exc()
        if 'conn' in locals() and conn:
            conn.rollback()
        raise HTTPException(status_code=500, detail=str(e))
    finally:
        if 'conn' in locals() and conn:
            conn.close()


# Note: All legacy programacion endpoints have been moved to app.modules.programacion.router.




# --- Roles & Permissions Endpoints (Using Supabase REST API) ---

def _get_supabase_headers():
    """Get headers for Supabase REST API calls"""
    service_key = os.getenv("SUPABASE_SERVICE_ROLE_KEY", "")
    return {
        "apikey": service_key,
        "Authorization": f"Bearer {service_key}",
        "Content-Type": "application/json",
        "Prefer": "return=representation"
    }

def _get_supabase_url():
    """Get Supabase REST API base URL"""
    url = os.getenv("SUPABASE_URL", "https://db.geofal.com.pe")
    return f"{url}/rest/v1"


def _permission(read: bool = False, write: bool = False, delete: bool = False) -> dict[str, bool]:
    return {"read": read, "write": write, "delete": delete}


_PERMISSION_MODULE_KEYS: tuple[str, ...] = (
    "tracing",
    "ingenieria_archivos",
    "clientes",
    "proyectos",
    "cotizadora",
    "programacion",
    "recepcion",
    "verificacion_muestras",
    "compresion",
    "control_probetas",
    "humedad",
    "cont_humedad",
    "cbr",
    "proctor",
    "llp",
    "gran_suelo",
    "gran_agregado",
    "cont_mat_organica",
    "terrones_fino_grueso",
    "azul_metileno",
    "part_livianas",
    "imp_organicas",
    "sul_magnesio",
    "angularidad",
    "abra",
    "abrass",
    "peso_unitario",
    "tamiz",
    "planas",
    "caras",
    "equi_arena",
    "ge_fino",
    "ge_grueso",
    "cd",
    "ph",
    "cloro_soluble",
    "sales_solubles",
    "sulfatos_solubles",
    "compresion_no_confinada",
    "laboratorio",
    "oficina_tecnica",
    "comercial",
    "administracion",
    "usuarios",
    "permisos",
    "auditoria",
    "configuracion",
)

_PERMISSION_KEY_ALIASES: dict[str, str] = {
    "correlativos": "ingenieria_archivos",
    "control_informes": "ingenieria_archivos",
    "verificacion": "verificacion_muestras",
}

_ROLE_ID_ALIASES: dict[str, str] = {
    "comercial": "auxiliar_comercial",
    "vendor": "auxiliar_comercial",
    "vendedor": "auxiliar_comercial",
    "sig_el_rol": "auxiliar_comercial",
    "tecnico_general": "tecnico",
    "tecnico_no_lab_write": "tecnico",
    "laboratorio_tipificador_no_lab_write": "laboratorio_lector",
    "oficina_tecnica_humedad": "oficina_tecnica",
    "oficina_tecnica_humedad_tipificador": "oficina_tecnica",
    "oficina_tecnica_sup": "oficina_tecnica",
}

_CONTROL_PERMISSION_MODULE_KEYS: tuple[str, ...] = (
    "ingenieria_archivos",
    "laboratorio",
    "oficina_tecnica",
    "comercial",
    "administracion",
)
_OFICINA_TECNICA_DELETE_MODULE_KEYS: tuple[str, ...] = (
    "verificacion",
    "verificacion_muestras",
    "recepcion",
    "compresion",
    "humedad",
    "cont_humedad",
    "planas",
    "caras",
    "cbr",
    "proctor",
    "llp",
    "gran_suelo",
    "gran_agregado",
    "cont_mat_organica",
    "terrones_fino_grueso",
    "azul_metileno",
    "part_livianas",
    "imp_organicas",
    "sul_magnesio",
    "angularidad",
    "abra",
    "abrass",
    "peso_unitario",
    "tamiz",
    "equi_arena",
    "ge_fino",
    "ge_grueso",
    "cd",
    "ph",
    "cloro_soluble",
    "sales_solubles",
    "sulfatos_solubles",
    "compresion_no_confinada",
)
_RESTRICTED_TECHNICAL_MODULE_KEYS: tuple[str, ...] = (
    "clientes",
    "proyectos",
    "cotizadora",
    "programacion",
)
_RESTRICTED_TECHNICAL_ROLE_IDS: set[str] = {"tecnico", "tecnico_suelos"}


def _is_restricted_technical_role(role_id: str | None) -> bool:
    normalized = (role_id or "").strip().lower()
    return normalized in _RESTRICTED_TECHNICAL_ROLE_IDS


def _available_modules_for_role(role_id: str | None) -> list[str]:
    modules = list(_PERMISSION_MODULE_KEYS)
    if _is_restricted_technical_role(role_id):
        return [module for module in modules if module not in _RESTRICTED_TECHNICAL_MODULE_KEYS]
    return modules


def _strip_control_permissions(permission_map: dict[str, dict[str, bool]] | None, role_id: str | None = None) -> dict[str, dict[str, bool]]:
    sanitized = dict(permission_map or {})
    normalized_role = _normalize_role_name(role_id)
    for module_key in _CONTROL_PERMISSION_MODULE_KEYS:
        if normalized_role == "tecnico_suelos" and module_key == "configuracion":
            continue
        sanitized[module_key] = _permission(False, False, False)
    for module_key in _RESTRICTED_TECHNICAL_MODULE_KEYS:
        sanitized[module_key] = _permission(False, False, False)
    if "correlativos" in sanitized:
        sanitized["correlativos"] = dict(sanitized["ingenieria_archivos"])
    return sanitized


def _sanitize_permissions_for_role(role_id: str | None, permission_map: dict[str, dict[str, bool]] | None) -> dict[str, dict[str, bool]]:
    sanitized = _strip_control_permissions(permission_map, role_id) if _is_restricted_technical_role(role_id) else dict(permission_map or {})
    if _normalize_role_name(role_id) == "tecnico_suelos":
        sanitized.setdefault("configuracion", _permission(True, False, False))
    return sanitized


def _grant_delete_to_oficina_tecnica(permission_map: dict[str, dict[str, bool]] | None, role_id: str | None) -> dict[str, dict[str, bool]]:
    normalized_role = _normalize_role_name(role_id)
    if not normalized_role.startswith("oficina_tecnica"):
        return dict(permission_map or {})

    granted = dict(permission_map or {})
    for module_key in _OFICINA_TECNICA_DELETE_MODULE_KEYS:
        module_permissions = granted.get(module_key)
        if isinstance(module_permissions, dict) and module_permissions.get("write") is True:
            granted[module_key] = _permission(
                bool(module_permissions.get("read")),
                True,
                True,
            )
    return granted


def _extra_special_lab_permissions(role_id: str) -> dict[str, dict[str, bool]]:
    normalized = (role_id or "").strip().lower()
    new_modules = (
        "cont_mat_organica",
        "terrones_fino_grueso",
        "azul_metileno",
        "part_livianas",
        "imp_organicas",
        "sul_magnesio",
        "angularidad",
    )

    if normalized == "admin":
        return {module: _permission(True, True, True) for module in new_modules}
    if normalized in {"laboratorio", "tecnico_suelos"}:
        return {module: _permission(True, True, False) for module in new_modules}
    return {module: _permission(False, False, False) for module in new_modules}


def _apply_role_permission_extensions(role_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    extended_roles: list[dict[str, Any]] = []
    for row in role_rows:
        role_data = dict(row)
        normalized_role = str(role_data.get("role_id") or "").strip().lower()
        permissions = dict(role_data.get("permissions") or {})
        extra_permissions = _extra_special_lab_permissions(str(role_data.get("role_id") or ""))
        for module_key, permission in extra_permissions.items():
            permissions.setdefault(module_key, permission)

        # Compatibility: keep old/new module keys aligned
        if "ingenieria_archivos" not in permissions and "correlativos" in permissions:
            permissions["ingenieria_archivos"] = permissions["correlativos"]
        if "correlativos" not in permissions and "ingenieria_archivos" in permissions:
            permissions["correlativos"] = permissions["ingenieria_archivos"]

        permissions = _sanitize_permissions_for_role(normalized_role, permissions)

        # Business rule: Administración must have access to Correlativo ING
        if normalized_role in {"administracion", "administrativo"}:
            permissions["ingenieria_archivos"] = _permission(True, True, False)
            permissions["correlativos"] = _permission(True, True, False)

        permissions = _grant_delete_to_oficina_tecnica(permissions, normalized_role)

        role_data["permissions"] = permissions
        extended_roles.append(role_data)
    return extended_roles


def _extract_request_user_id(request: Request) -> str | None:
    user_payload = getattr(request.state, "user", None)
    if not isinstance(user_payload, dict):
        return None
    candidate = user_payload.get("sub") or user_payload.get("id")
    return str(candidate).strip() if candidate else None


def _normalize_role_name(value: str | None) -> str:
    normalized = (value or "").strip().lower()
    return _ROLE_ID_ALIASES.get(normalized, normalized)


def _canonicalize_role_definition_rows(role_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    canonical_rows: dict[str, dict[str, Any]] = {}
    for row in role_rows:
        role_data = dict(row)
        raw_role_id = str(role_data.get("role_id") or "").strip().lower()
        canonical_role_id = _normalize_role_name(raw_role_id)
        role_data["role_id"] = canonical_role_id
        role_data["_is_alias"] = canonical_role_id != raw_role_id

        current = canonical_rows.get(canonical_role_id)
        if current is None:
            canonical_rows[canonical_role_id] = role_data
            continue

        if current.get("_is_alias") and not role_data.get("_is_alias"):
            canonical_rows[canonical_role_id] = role_data

    return [
        {key: value for key, value in row.items() if key != "_is_alias"}
        for row in canonical_rows.values()
    ]


def _permission_from_payload(raw: Any) -> dict[str, bool]:
    source = raw if isinstance(raw, dict) else {}
    return {
        "read": bool(source.get("read", False)),
        "write": bool(source.get("write", False)),
        "delete": bool(source.get("delete", False)),
    }


def _normalize_permission_map(raw: Any) -> dict[str, dict[str, bool]]:
    if not isinstance(raw, dict):
        return {}

    normalized: dict[str, dict[str, bool]] = {}
    for module_key, permission_data in raw.items():
        key = str(module_key or "").strip().lower()
        if not key:
            continue
        key = _PERMISSION_KEY_ALIASES.get(key, key)
        if key not in _PERMISSION_MODULE_KEYS:
            continue
        normalized[key] = _permission_from_payload(permission_data)

    # Compatibility for older frontends still expecting "correlativos" key.
    if "ingenieria_archivos" in normalized and "correlativos" not in normalized:
        normalized["correlativos"] = dict(normalized["ingenieria_archivos"])
    return normalized


def _compact_permission_override(
    base_permissions: dict[str, dict[str, bool]] | None,
    override_permissions: dict[str, dict[str, bool]] | None,
) -> dict[str, dict[str, bool]]:
    base = _normalize_permission_map(base_permissions)
    override = _normalize_permission_map(override_permissions)
    compacted: dict[str, dict[str, bool]] = {}

    for module_key, permission in override.items():
        base_permission = base.get(module_key, _permission(False, False, False))
        if permission != base_permission:
            compacted[module_key] = permission

    return compacted


def _merge_permission_maps(
    base: dict[str, dict[str, bool]] | None,
    override: dict[str, dict[str, bool]] | None,
) -> dict[str, dict[str, bool]]:
    merged: dict[str, dict[str, bool]] = {}
    for module in _PERMISSION_MODULE_KEYS:
        merged[module] = _permission_from_payload((base or {}).get(module))
    for module, values in (override or {}).items():
        canonical = _PERMISSION_KEY_ALIASES.get(module, module)
        if canonical in _PERMISSION_MODULE_KEYS:
            merged[canonical] = _permission_from_payload(values)
    if "ingenieria_archivos" in merged:
        merged["correlativos"] = dict(merged["ingenieria_archivos"])
    return merged


def _get_profile_role(cur, user_id: str) -> str | None:
    cur.execute("SELECT role FROM perfiles WHERE id = %s LIMIT 1", (user_id,))
    row = cur.fetchone()
    if not row:
        return None
    if isinstance(row, dict):
        return row.get("role")
    return row[0] if row else None


def _resolve_profile_avatar_url(cur, user_id: str | None) -> str | None:
    normalized_user_id = str(user_id or "").strip()
    if not normalized_user_id:
        return None
    cur.execute("SELECT avatar_url FROM perfiles WHERE id = %s LIMIT 1", (normalized_user_id,))
    row = cur.fetchone()
    if not row:
        return None
    if isinstance(row, dict):
        avatar_url = str(row.get("avatar_url") or "").strip()
    else:
        avatar_url = str(row[0] or "").strip()
    return avatar_url or None


def _notification_identity_key(user_id: str, role_id: str, module_key: str = "laboratorio") -> str:
    return f"{user_id}:{role_id}:{module_key}"


def _permission_read_write(permission_map: dict[str, dict[str, bool]] | None, module_key: str) -> tuple[bool, bool]:
    module = (permission_map or {}).get(module_key) or {}
    return module.get("read") is True, module.get("write") is True


def _build_permission_conflict_notifications(cur) -> list[dict[str, Any]]:
    cur.execute(
        """
        SELECT
            p.id::text AS user_id,
            p.full_name,
            p.email,
            p.role,
            rd.permissions AS role_permissions,
            up.enabled AS override_enabled,
            up.permissions AS override_permissions,
            COALESCE(up.updated_at, rd.updated_at, NOW()) AS updated_at
        FROM perfiles p
        LEFT JOIN role_definitions rd ON rd.role_id = p.role
        LEFT JOIN user_permission_overrides up ON up.user_id = p.id
        WHERE p.role IN ('jefe_laboratorio', 'laboratorio_tipificador')
        ORDER BY p.full_name ASC
        """
    )
    rows = cur.fetchall() or []
    notifications: list[dict[str, Any]] = []

    for row in rows:
        if not isinstance(row, dict):
            continue
        role_id = _normalize_role_name(row.get("role"))
        base_permissions = _normalize_permission_map(row.get("role_permissions"))
        override_enabled = bool(row.get("override_enabled"))
        override_permissions = _normalize_permission_map(row.get("override_permissions"))
        effective_permissions = _merge_permission_maps(base_permissions, override_permissions if override_enabled else {})

        base_read, base_write = _permission_read_write(base_permissions, "laboratorio")
        effective_read, effective_write = _permission_read_write(effective_permissions, "laboratorio")
        override_read, override_write = _permission_read_write(override_permissions, "laboratorio")

        if effective_write is True:
            continue

        reason = "override_granular" if override_enabled and override_write is False and base_write is True else "role_definition"
        title = "Permiso de laboratorio inconsistente"
        message = (
            f"{row.get('full_name') or row.get('email') or row.get('user_id')} quedó sin edición en Laboratorio "
            f"para el rol {role_id}."
        )
        if reason == "override_granular":
            message += " Un override granular parece haber desactivado la escritura."
        elif not base_write:
            message += " La definición base del rol no otorga escritura."

        notifications.append({
            "id": _notification_identity_key(str(row.get("user_id") or ""), role_id),
            "type": "permission_conflict",
            "severity": "warning",
            "title": title,
            "message": message,
            "created_at": row.get("updated_at") or datetime.utcnow(),
            "metadata": {
                "user_id": row.get("user_id"),
                "full_name": row.get("full_name"),
                "email": row.get("email"),
                "role": role_id,
                "module": "laboratorio",
                "base_read": base_read,
                "base_write": base_write,
                "effective_read": effective_read,
                "effective_write": effective_write,
                "override_enabled": override_enabled,
                "override_write": override_write,
                "reason": reason,
            },
        })

    return notifications


def _ensure_dashboard_notifications_table(cur) -> None:
    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS dashboard_notifications (
            notification_key TEXT PRIMARY KEY,
            type TEXT NOT NULL,
            severity TEXT NOT NULL DEFAULT 'warning',
            title TEXT NOT NULL,
            message TEXT NOT NULL,
            status TEXT NOT NULL DEFAULT 'open',
            metadata JSONB NOT NULL DEFAULT '{}'::jsonb,
            created_at TIMESTAMPTZ NOT NULL DEFAULT NOW(),
            updated_at TIMESTAMPTZ NOT NULL DEFAULT NOW(),
            acknowledged_at TIMESTAMPTZ NULL,
            acknowledged_by UUID NULL,
            resolved_at TIMESTAMPTZ NULL,
            resolved_by UUID NULL,
            last_detected_at TIMESTAMPTZ NOT NULL DEFAULT NOW(),
            CONSTRAINT dashboard_notifications_status_check
                CHECK (status IN ('open', 'acknowledged', 'resolved'))
        );
        """
    )
    cur.execute(
        """
        CREATE INDEX IF NOT EXISTS idx_dashboard_notifications_status
        ON dashboard_notifications (status);
        """
    )
    cur.execute(
        """
        CREATE INDEX IF NOT EXISTS idx_dashboard_notifications_type_status
        ON dashboard_notifications (type, status);
        """
    )
    cur.execute(
        """
        CREATE INDEX IF NOT EXISTS idx_dashboard_notifications_last_detected_at
        ON dashboard_notifications (last_detected_at DESC);
        """
    )


def _upsert_dashboard_notification(cur, notification: dict[str, Any]) -> None:
    metadata = notification.get("metadata") or {}
    cur.execute(
        """
        INSERT INTO dashboard_notifications (
            notification_key,
            type,
            severity,
            title,
            message,
            status,
            metadata,
            created_at,
            updated_at,
            acknowledged_at,
            acknowledged_by,
            resolved_at,
            resolved_by,
            last_detected_at
        )
        VALUES (
            %s,
            %s,
            %s,
            %s,
            %s,
            'open',
            %s::jsonb,
            %s,
            NOW(),
            NULL,
            NULL,
            NULL,
            NULL,
            NOW()
        )
        ON CONFLICT (notification_key) DO UPDATE SET
            type = EXCLUDED.type,
            severity = EXCLUDED.severity,
            title = EXCLUDED.title,
            message = EXCLUDED.message,
            metadata = EXCLUDED.metadata,
            updated_at = NOW(),
            last_detected_at = NOW(),
            status = CASE
                WHEN dashboard_notifications.status = 'acknowledged' THEN 'acknowledged'
                ELSE 'open'
            END,
            acknowledged_at = CASE
                WHEN dashboard_notifications.status = 'acknowledged' THEN dashboard_notifications.acknowledged_at
                ELSE NULL
            END,
            acknowledged_by = CASE
                WHEN dashboard_notifications.status = 'acknowledged' THEN dashboard_notifications.acknowledged_by
                ELSE NULL
            END,
            resolved_at = NULL,
            resolved_by = NULL
        """,
        (
            str(notification.get("id") or ""),
            str(notification.get("type") or "permission_conflict"),
            str(notification.get("severity") or "warning"),
            str(notification.get("title") or "Notificación"),
            str(notification.get("message") or ""),
            Json(metadata),
            notification.get("created_at") or datetime.utcnow(),
        ),
    )


def _mark_resolved_dashboard_notifications(cur, current_user_id: str, active_keys: list[str]) -> None:
    if active_keys:
        cur.execute(
            """
            UPDATE dashboard_notifications
            SET
                status = 'resolved',
                resolved_at = NOW(),
                resolved_by = %s::uuid,
                updated_at = NOW()
            WHERE type = 'permission_conflict'
              AND status IN ('open', 'acknowledged')
              AND NOT (notification_key = ANY(%s))
            """,
            (current_user_id, active_keys),
        )
        return

    cur.execute(
        """
        UPDATE dashboard_notifications
        SET
            status = 'resolved',
            resolved_at = NOW(),
            resolved_by = %s::uuid,
            updated_at = NOW()
        WHERE type = 'permission_conflict'
          AND status IN ('open', 'acknowledged')
        """,
        (current_user_id,),
    )


def _fetch_dashboard_notifications(cur) -> list[dict[str, Any]]:
    cur.execute(
        """
        SELECT
            notification_key AS id,
            type,
            severity,
            title,
            message,
            status,
            created_at,
            acknowledged_at,
            metadata
        FROM dashboard_notifications
        WHERE type = 'permission_conflict'
          AND status IN ('open', 'acknowledged')
        ORDER BY
            CASE WHEN status = 'open' THEN 0 ELSE 1 END,
            last_detected_at DESC,
            created_at DESC,
            notification_key ASC
        """
    )
    rows = cur.fetchall() or []
    notifications: list[dict[str, Any]] = []
    for row in rows:
        if isinstance(row, dict):
            notifications.append(dict(row))
        elif row:
            notifications.append({
                "id": row[0],
                "type": row[1],
                "severity": row[2],
                "title": row[3],
                "message": row[4],
                "status": row[5],
                "created_at": row[6],
                "acknowledged_at": row[7] if len(row) > 7 else None,
                "metadata": row[8] if len(row) > 8 else {},
            })
    return notifications


def _fetch_dashboard_notification_history(cur, limit: int = 12) -> list[dict[str, Any]]:
    cur.execute(
        """
        SELECT
            notification_key AS id,
            type,
            severity,
            title,
            message,
            status,
            created_at,
            acknowledged_at,
            resolved_at,
            metadata
        FROM dashboard_notifications
        WHERE type = 'permission_conflict'
          AND status = 'resolved'
        ORDER BY resolved_at DESC NULLS LAST, last_detected_at DESC, created_at DESC
        LIMIT %s
        """,
        (limit,),
    )
    rows = cur.fetchall() or []
    history: list[dict[str, Any]] = []
    for row in rows:
        if isinstance(row, dict):
            history.append(dict(row))
        elif row:
            history.append({
                "id": row[0],
                "type": row[1],
                "severity": row[2],
                "title": row[3],
                "message": row[4],
                "status": row[5],
                "created_at": row[6],
                "acknowledged_at": row[7] if len(row) > 7 else None,
                "resolved_at": row[8] if len(row) > 8 else None,
                "metadata": row[9] if len(row) > 9 else {},
            })
    return history


def _fetch_quote_notifications(cur, role_id: str, limit: int = 12) -> list[dict[str, Any]]:
    safe_limit = max(1, min(int(limit or 12), 100))
    cur.execute(
        """
        SELECT
            notification_key AS id,
            type,
            severity,
            title,
            message,
            status,
            created_at,
            updated_at,
            metadata
        FROM dashboard_notifications
        WHERE type = 'quote_created'
          AND COALESCE(metadata->'audience_roles', '[]'::jsonb) ? %s
        ORDER BY last_detected_at DESC, created_at DESC
        LIMIT %s
        """,
        (role_id, safe_limit),
    )
    rows = cur.fetchall() or []
    notifications: list[dict[str, Any]] = []
    for row in rows:
        if isinstance(row, dict):
            notification = dict(row)
            metadata = dict(notification.get("metadata") or {})
            metadata.setdefault("created_by_avatar_url", _resolve_profile_avatar_url(cur, metadata.get("created_by_user_id")))
            notification["metadata"] = metadata
            notifications.append(notification)
        elif row:
            metadata = row[8] if len(row) > 8 else {}
            metadata = dict(metadata or {})
            metadata.setdefault("created_by_avatar_url", _resolve_profile_avatar_url(cur, metadata.get("created_by_user_id")))
            notifications.append({
                "id": row[0],
                "type": row[1],
                "severity": row[2],
                "title": row[3],
                "message": row[4],
                "status": row[5],
                "created_at": row[6],
                "updated_at": row[7] if len(row) > 7 else None,
                "metadata": metadata,
            })
    return notifications


def _fetch_laboratory_notifications(cur, role_id: str, limit: int = 12) -> list[dict[str, Any]]:
    safe_limit = max(1, min(int(limit or 12), 100))
    cur.execute(
        """
        SELECT
            notification_key AS id,
            type,
            severity,
            title,
            message,
            status,
            created_at,
            updated_at,
            metadata
        FROM dashboard_notifications
        WHERE type IN ('lab_essay_created', 'lab_essay_updated', 'lab_essay_deleted')
          AND COALESCE(metadata->'audience_roles', '[]'::jsonb) ? %s
        ORDER BY last_detected_at DESC, created_at DESC
        LIMIT %s
        """,
        (role_id, safe_limit),
    )
    rows = cur.fetchall() or []
    notifications: list[dict[str, Any]] = []
    for row in rows:
        metadata = row.get("metadata") or {}
        avatar_url = metadata.get("created_by_avatar_url") or _resolve_profile_avatar_url(cur, metadata.get("created_by_user_id"))
        notifications.append(
            {
                "id": row.get("id"),
                "type": row.get("type") or "lab_essay_created",
                "severity": row.get("severity") or "info",
                "title": row.get("title") or "Ensayo de laboratorio",
                "message": row.get("message") or "",
                "status": row.get("status") or "open",
                "created_at": row.get("created_at"),
                "updated_at": row.get("updated_at"),
                "metadata": {
                    **metadata,
                    "created_by_avatar_url": avatar_url,
                    "module": metadata.get("module"),
                    "module_label": metadata.get("module_label"),
                    "record_id": metadata.get("record_id"),
                    "record_code": metadata.get("record_code"),
                    "action": metadata.get("action"),
                    "created_by": metadata.get("created_by"),
                    "created_by_user_id": metadata.get("created_by_user_id"),
                    "created_by_role": metadata.get("created_by_role"),
                    "audience_roles": metadata.get("audience_roles") or [],
                    "detail_module": metadata.get("detail_module"),
                    "detail_record_id": metadata.get("detail_record_id"),
                },
            }
        )
    return notifications


def _sync_dashboard_notifications(cur, current_user_id: str) -> list[dict[str, Any]]:
    _ensure_dashboard_notifications_table(cur)
    derived_notifications = _build_permission_conflict_notifications(cur)
    active_keys: list[str] = []

    for notification in derived_notifications:
        notification_key = str(notification.get("id") or "").strip()
        if not notification_key:
            continue
        active_keys.append(notification_key)
        _upsert_dashboard_notification(cur, notification)

    _mark_resolved_dashboard_notifications(cur, current_user_id, active_keys)
    return _fetch_dashboard_notifications(cur)


def _acknowledge_dashboard_notification(cur, notification_key: str, current_user_id: str) -> dict[str, Any] | None:
    _ensure_dashboard_notifications_table(cur)
    cur.execute(
        """
        UPDATE dashboard_notifications
        SET
            status = 'acknowledged',
            acknowledged_at = COALESCE(acknowledged_at, NOW()),
            acknowledged_by = %s::uuid,
            updated_at = NOW()
        WHERE notification_key = %s
          AND status IN ('open', 'acknowledged')
        RETURNING
            notification_key AS id,
            type,
            severity,
            title,
            message,
            status,
            created_at,
            acknowledged_at,
            metadata
        """,
        (current_user_id, notification_key),
    )
    row = cur.fetchone()
    if not row:
        return None
    if isinstance(row, dict):
        return dict(row)
    return {
        "id": row[0],
        "type": row[1],
        "severity": row[2],
        "title": row[3],
        "message": row[4],
        "status": row[5],
        "created_at": row[6],
        "acknowledged_at": row[7] if len(row) > 7 else None,
        "metadata": row[8] if len(row) > 8 else {},
    }


def _get_dashboard_notification_by_key(cur, notification_key: str) -> dict[str, Any] | None:
    _ensure_dashboard_notifications_table(cur)
    cur.execute(
        """
        SELECT
            notification_key AS id,
            type,
            severity,
            title,
            message,
            status,
            created_at,
            acknowledged_at,
            metadata
        FROM dashboard_notifications
        WHERE notification_key = %s
          AND status IN ('open', 'acknowledged')
        LIMIT 1
        """,
        (notification_key,),
    )
    row = cur.fetchone()
    if not row:
        return None
    if isinstance(row, dict):
        return dict(row)
    return {
        "id": row[0],
        "type": row[1],
        "severity": row[2],
        "title": row[3],
        "message": row[4],
        "status": row[5],
        "created_at": row[6],
        "acknowledged_at": row[7] if len(row) > 7 else None,
        "metadata": row[8] if len(row) > 8 else {},
    }


def _notification_audience_roles(notification: dict[str, Any]) -> set[str]:
    metadata = notification.get("metadata") or {}
    audience_roles = metadata.get("audience_roles") or []
    return {str(role or "").strip().lower() for role in audience_roles if str(role or "").strip()}


def _can_acknowledge_notification(role: str, notification: dict[str, Any]) -> bool:
    normalized_role = _normalize_role_name(role)
    notification_type = str(notification.get("type") or "")
    audience_roles = _notification_audience_roles(notification)

    if normalized_role in {"admin", "admin_general"}:
        return notification_type == "permission_conflict"

    if normalized_role == "auxiliar_comercial":
        return notification_type == "quote_created" and "auxiliar_comercial" in audience_roles

    if normalized_role in {"jefe_laboratorio", "laboratorio_tipificador"}:
        return notification_type in {"lab_essay_created", "lab_essay_updated", "lab_essay_deleted"} and bool(
            audience_roles.intersection({"jefe_laboratorio", "laboratorio_tipificador"})
        )

    return False


@app.get("/roles")
async def get_roles():
    """Get all role definitions using Supabase REST API"""
    try:
        url = f"{_get_supabase_url()}/role_definitions?order=label.asc"
        logger.info(f"[Auth] Fetching role definitions from: {url}")
        response = http_get(
            url,
            headers=_get_supabase_headers(),
            timeout=3,
            request_name="supabase.role_definitions.list",
        )
        
        if response.status_code == 404 or (response.status_code == 200 and response.json() == []):
            # Table doesn't exist or is empty - return default roles
            return _apply_role_permission_extensions([
                {
                    "role_id": "admin",
                    "label": "Administrador",
                    "description": "Acceso completo al sistema",
                    "permissions": {
                        "clientes": {"read": True, "write": True, "delete": True},
                        "proyectos": {"read": True, "write": True, "delete": True},
                        "cotizadora": {"read": True, "write": True, "delete": True},
                        "programacion": {"read": True, "write": True, "delete": True},
                        "recepcion": {"read": True, "write": True, "delete": True},
                        "verificacion_muestras": {"read": True, "write": True, "delete": True},
                        "compresion": {"read": True, "write": True, "delete": True},
                        "tracing": {"read": True, "write": True, "delete": True},
                        "control_probetas": {"read": True, "write": True, "delete": True},
                        "humedad": {"read": True, "write": True, "delete": True},
                        "cont_humedad": {"read": True, "write": True, "delete": True},
                        "planas": {"read": True, "write": True, "delete": True},
                        "caras": {"read": True, "write": True, "delete": True},
                        "cbr": {"read": True, "write": True, "delete": True},
                        "proctor": {"read": True, "write": True, "delete": True},
                        "llp": {"read": True, "write": True, "delete": True},
                        "gran_suelo": {"read": True, "write": True, "delete": True},
                        "gran_agregado": {"read": True, "write": True, "delete": True},
                        "abra": {"read": True, "write": True, "delete": True},
                        "abrass": {"read": True, "write": True, "delete": True},
                        "peso_unitario": {"read": True, "write": True, "delete": True},
                        "tamiz": {"read": True, "write": True, "delete": True},
                        "equi_arena": {"read": True, "write": True, "delete": True},
                        "ge_fino": {"read": True, "write": True, "delete": True},
                        "ge_grueso": {"read": True, "write": True, "delete": True},
                        "cd": {"read": True, "write": True, "delete": True},
                        "ph": {"read": True, "write": True, "delete": True},
                        "cloro_soluble": {"read": True, "write": True, "delete": True},
                        "sales_solubles": {"read": True, "write": True, "delete": True},
                        "sulfatos_solubles": {"read": True, "write": True, "delete": True},
                        "compresion_no_confinada": {"read": True, "write": True, "delete": True},
                        "usuarios": {"read": True, "write": True, "delete": True},
                        "auditoria": {"read": True, "write": True, "delete": True},
                        "configuracion": {"read": True, "write": True, "delete": True},
                        "laboratorio": {"read": True, "write": True, "delete": True},
                        "comercial": {"read": True, "write": True, "delete": True},
                        "administracion": {"read": True, "write": True, "delete": True},
                        "permisos": {"read": True, "write": True, "delete": True},
                        "correlativos": {"read": True, "write": True, "delete": True}
                    },
                    "is_system": True
                },
                {
                    "role_id": "auxiliar_comercial",
                    "label": "Auxiliar Comercial",
                    "description": "Soporte comercial (clientes, proyectos y cotizaciones)",
                    "permissions": {
                        "clientes": {"read": True, "write": True, "delete": False},
                        "proyectos": {"read": True, "write": True, "delete": False},
                        "cotizadora": {"read": True, "write": True, "delete": False},
                        "programacion": {"read": True, "write": False, "delete": False},
                        "laboratorio": {"read": True, "write": False, "delete": False},
                        "recepcion": {"read": False, "write": False, "delete": False},
                        "verificacion_muestras": {"read": False, "write": False, "delete": False},
                        "compresion": {"read": False, "write": False, "delete": False},
                        "tracing": {"read": False, "write": False, "delete": False},
                        "humedad": {"read": False, "write": False, "delete": False},
                        "cont_humedad": {"read": False, "write": False, "delete": False},
                        "planas": {"read": False, "write": False, "delete": False},
                        "caras": {"read": False, "write": False, "delete": False},
                        "cbr": {"read": False, "write": False, "delete": False},
                        "proctor": {"read": False, "write": False, "delete": False},
                        "llp": {"read": False, "write": False, "delete": False},
                        "gran_suelo": {"read": False, "write": False, "delete": False},
                        "gran_agregado": {"read": False, "write": False, "delete": False},
                        "abra": {"read": False, "write": False, "delete": False},
                        "abrass": {"read": False, "write": False, "delete": False},
                        "peso_unitario": {"read": False, "write": False, "delete": False},
                        "tamiz": {"read": False, "write": False, "delete": False},
                        "equi_arena": {"read": False, "write": False, "delete": False},
                        "ge_fino": {"read": False, "write": False, "delete": False},
                        "ge_grueso": {"read": False, "write": False, "delete": False},
                        "cd": {"read": False, "write": False, "delete": False},
                        "ph": {"read": False, "write": False, "delete": False},
                        "cloro_soluble": {"read": False, "write": False, "delete": False},
                        "sales_solubles": {"read": False, "write": False, "delete": False},
                        "sulfatos_solubles": {"read": False, "write": False, "delete": False},
                        "compresion_no_confinada": {"read": False, "write": False, "delete": False},
                        "usuarios": {"read": False, "write": False, "delete": False},
                        "auditoria": {"read": False, "write": False, "delete": False},
                        "configuracion": {"read": False, "write": False, "delete": False},
                        "laboratorio": {"read": False, "write": False, "delete": False},
                        "comercial": {"read": True, "write": True, "delete": False},
                        "administracion": {"read": False, "write": False, "delete": False},
                        "permisos": {"read": False, "write": False, "delete": False},
                        "correlativos": {"read": False, "write": False, "delete": False}
                    },
                    "is_system": True
                },
                {
                    "role_id": "tecnico",
                    "label": "Técnico",
                    "description": "Acceso técnico a ensayos sin clientes, proyectos, programación ni tablas de control",
                    "permissions": {
                        "recepcion": {"read": True, "write": True, "delete": True},
                        "verificacion_muestras": {"read": True, "write": True, "delete": True},
                        "compresion": {"read": True, "write": True, "delete": True},
                        "tracing": {"read": True, "write": True, "delete": True},
                        "humedad": {"read": True, "write": True, "delete": True},
                        "cont_humedad": {"read": True, "write": True, "delete": True},
                        "planas": {"read": True, "write": True, "delete": True},
                        "caras": {"read": True, "write": True, "delete": True},
                        "cbr": {"read": True, "write": True, "delete": True},
                        "proctor": {"read": True, "write": True, "delete": True},
                        "llp": {"read": True, "write": True, "delete": True},
                        "gran_suelo": {"read": True, "write": True, "delete": True},
                        "gran_agregado": {"read": True, "write": True, "delete": True},
                        "cont_mat_organica": {"read": True, "write": True, "delete": True},
                        "terrones_fino_grueso": {"read": True, "write": True, "delete": True},
                        "azul_metileno": {"read": True, "write": True, "delete": True},
                        "part_livianas": {"read": True, "write": True, "delete": True},
                        "imp_organicas": {"read": True, "write": True, "delete": True},
                        "sul_magnesio": {"read": True, "write": True, "delete": True},
                        "angularidad": {"read": True, "write": True, "delete": True},
                        "abra": {"read": True, "write": True, "delete": True},
                        "abrass": {"read": True, "write": True, "delete": True},
                        "peso_unitario": {"read": True, "write": True, "delete": True},
                        "tamiz": {"read": True, "write": True, "delete": True},
                        "equi_arena": {"read": True, "write": True, "delete": True},
                        "ge_fino": {"read": True, "write": True, "delete": True},
                        "ge_grueso": {"read": True, "write": True, "delete": True},
                        "cd": {"read": True, "write": True, "delete": True},
                        "ph": {"read": True, "write": True, "delete": True},
                        "cloro_soluble": {"read": True, "write": True, "delete": True},
                        "sales_solubles": {"read": True, "write": True, "delete": True},
                        "sulfatos_solubles": {"read": True, "write": True, "delete": True},
                        "compresion_no_confinada": {"read": True, "write": True, "delete": True},
                        "configuracion": {"read": True, "write": False, "delete": False},
                        "usuarios": {"read": False, "write": False, "delete": False},
                        "auditoria": {"read": False, "write": False, "delete": False},
                        "comercial": {"read": False, "write": False, "delete": False},
                        "administracion": {"read": False, "write": False, "delete": False},
                        "permisos": {"read": False, "write": False, "delete": False},
                        "correlativos": {"read": False, "write": False, "delete": False}
                    },
                    "is_system": True
                },
                {
                    "role_id": "laboratorio_tipificador",
                    "label": "Laboratorio Tipificador",
                    "description": "Acceso técnico a ensayos con tabla de control lab en edición",
                    "permissions": {
                        "recepcion": {"read": True, "write": True, "delete": False},
                        "verificacion_muestras": {"read": True, "write": True, "delete": False},
                        "compresion": {"read": True, "write": True, "delete": False},
                        "tracing": {"read": True, "write": True, "delete": False},
                        "humedad": {"read": True, "write": True, "delete": False},
                        "cont_humedad": {"read": True, "write": True, "delete": False},
                        "planas": {"read": True, "write": True, "delete": False},
                        "caras": {"read": True, "write": True, "delete": False},
                        "cbr": {"read": True, "write": True, "delete": False},
                        "proctor": {"read": True, "write": True, "delete": False},
                        "llp": {"read": True, "write": True, "delete": False},
                        "gran_suelo": {"read": True, "write": True, "delete": False},
                        "gran_agregado": {"read": True, "write": True, "delete": False},
                        "cont_mat_organica": {"read": True, "write": True, "delete": False},
                        "terrones_fino_grueso": {"read": True, "write": True, "delete": False},
                        "azul_metileno": {"read": True, "write": True, "delete": False},
                        "part_livianas": {"read": True, "write": True, "delete": False},
                        "imp_organicas": {"read": True, "write": True, "delete": False},
                        "sul_magnesio": {"read": True, "write": True, "delete": False},
                        "angularidad": {"read": True, "write": True, "delete": False},
                        "abra": {"read": True, "write": True, "delete": False},
                        "abrass": {"read": True, "write": True, "delete": False},
                        "peso_unitario": {"read": True, "write": True, "delete": False},
                        "tamiz": {"read": True, "write": True, "delete": False},
                        "equi_arena": {"read": True, "write": True, "delete": False},
                        "ge_fino": {"read": True, "write": True, "delete": False},
                        "ge_grueso": {"read": True, "write": True, "delete": False},
                        "cd": {"read": True, "write": True, "delete": False},
                        "ph": {"read": True, "write": True, "delete": False},
                        "cloro_soluble": {"read": True, "write": True, "delete": False},
                        "sales_solubles": {"read": True, "write": True, "delete": False},
                        "sulfatos_solubles": {"read": True, "write": True, "delete": False},
                        "compresion_no_confinada": {"read": True, "write": True, "delete": False},
                        "laboratorio": {"read": True, "write": True, "delete": False},
                        "configuracion": {"read": True, "write": False, "delete": False},
                        "usuarios": {"read": False, "write": False, "delete": False},
                        "auditoria": {"read": False, "write": False, "delete": False},
                        "comercial": {"read": False, "write": False, "delete": False},
                        "administracion": {"read": False, "write": False, "delete": False},
                        "permisos": {"read": False, "write": False, "delete": False},
                        "correlativos": {"read": False, "write": False, "delete": False}
                    },
                    "is_system": True
                },
                {
                    "role_id": "laboratorio_lector",
                    "label": "Lector Laboratorio",
                    "description": "Acceso técnico a ensayos con tabla de control lab en solo lectura",
                    "permissions": {
                        "recepcion": {"read": True, "write": True, "delete": False},
                        "verificacion_muestras": {"read": True, "write": True, "delete": False},
                        "compresion": {"read": True, "write": True, "delete": False},
                        "tracing": {"read": True, "write": True, "delete": False},
                        "humedad": {"read": True, "write": True, "delete": False},
                        "cont_humedad": {"read": True, "write": True, "delete": False},
                        "planas": {"read": True, "write": True, "delete": False},
                        "caras": {"read": True, "write": True, "delete": False},
                        "cbr": {"read": True, "write": True, "delete": False},
                        "proctor": {"read": True, "write": True, "delete": False},
                        "llp": {"read": True, "write": True, "delete": False},
                        "gran_suelo": {"read": True, "write": True, "delete": False},
                        "gran_agregado": {"read": True, "write": True, "delete": False},
                        "cont_mat_organica": {"read": True, "write": True, "delete": False},
                        "terrones_fino_grueso": {"read": True, "write": True, "delete": False},
                        "azul_metileno": {"read": True, "write": True, "delete": False},
                        "part_livianas": {"read": True, "write": True, "delete": False},
                        "imp_organicas": {"read": True, "write": True, "delete": False},
                        "sul_magnesio": {"read": True, "write": True, "delete": False},
                        "angularidad": {"read": True, "write": True, "delete": False},
                        "abra": {"read": True, "write": True, "delete": False},
                        "abrass": {"read": True, "write": True, "delete": False},
                        "peso_unitario": {"read": True, "write": True, "delete": False},
                        "tamiz": {"read": True, "write": True, "delete": False},
                        "equi_arena": {"read": True, "write": True, "delete": False},
                        "ge_fino": {"read": True, "write": True, "delete": False},
                        "ge_grueso": {"read": True, "write": True, "delete": False},
                        "cd": {"read": True, "write": True, "delete": False},
                        "ph": {"read": True, "write": True, "delete": False},
                        "cloro_soluble": {"read": True, "write": True, "delete": False},
                        "sales_solubles": {"read": True, "write": True, "delete": False},
                        "sulfatos_solubles": {"read": True, "write": True, "delete": False},
                        "compresion_no_confinada": {"read": True, "write": True, "delete": False},
                        "laboratorio": {"read": True, "write": False, "delete": False},
                        "configuracion": {"read": True, "write": False, "delete": False},
                        "usuarios": {"read": False, "write": False, "delete": False},
                        "auditoria": {"read": False, "write": False, "delete": False},
                        "comercial": {"read": False, "write": False, "delete": False},
                        "administracion": {"read": False, "write": False, "delete": False},
                        "permisos": {"read": False, "write": False, "delete": False},
                        "correlativos": {"read": False, "write": False, "delete": False}
                    },
                    "is_system": True
                }
            ])
        
        if response.status_code != 200:
            logger.error("Supabase error fetching roles: %s - %s", response.status_code, response.text)
            raise HTTPException(status_code=500, detail=f"Error fetching roles: {response.text}")

        return _apply_role_permission_extensions(_canonicalize_role_definition_rows(response.json()))
    except requests.RequestException as e:
        logger.exception("Request error fetching roles")
        # Return default roles on connection error
        return _apply_role_permission_extensions([
            {"role_id": "admin", "label": "Administrador", "description": "Acceso completo", "permissions": {}, "is_system": True},
            {"role_id": "auxiliar_comercial", "label": "Auxiliar Comercial", "description": "Acceso comercial", "permissions": {}, "is_system": True}
        ])


@app.get("/notifications")
async def get_notifications(request: Request):
    """Return admin dashboard notifications derived from permission inconsistencies."""
    if not _has_database_url():
        return {"data": [], "count": 0}

    current_user_id = _extract_request_user_id(request)
    if not current_user_id:
        raise HTTPException(status_code=401, detail="Usuario no autenticado")

    try:
        conn = _get_connection()
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            current_role = _normalize_role_name(_get_profile_role(cur, current_user_id))
            if current_role not in {"admin", "admin_general"}:
                return {"data": [], "count": 0}

            notifications = _sync_dashboard_notifications(cur, current_user_id)
            conn.commit()
            open_count = sum(1 for item in notifications if _normalize_role_name(str(item.get("status") or "")) == "open")
            return {"data": notifications, "count": open_count}
    except Exception as e:
        logger.warning("Error fetching notifications: %s", e)
        return {"data": [], "count": 0}
    finally:
        if 'conn' in locals() and conn:
            conn.close()


@app.get("/notifications/feed")
async def get_notifications_feed(request: Request, limit: int = 12):
    """Return the notification feed relevant to the current user's role."""
    if not _has_database_url():
        return {"data": [], "count": 0}

    current_user_id = _extract_request_user_id(request)
    if not current_user_id:
        raise HTTPException(status_code=401, detail="Usuario no autenticado")

    safe_limit = max(1, min(int(limit or 12), 100))

    try:
        conn = _get_connection()
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            current_role = _normalize_role_name(_get_profile_role(cur, current_user_id))

            if current_role in {"admin", "admin_general"}:
                notifications = _sync_dashboard_notifications(cur, current_user_id)
                conn.commit()
                open_count = sum(1 for item in notifications if _normalize_role_name(str(item.get("status") or "")) == "open")
                return {"data": notifications, "count": open_count}

            if current_role == "auxiliar_comercial":
                notifications = _fetch_quote_notifications(cur, current_role, limit=safe_limit)
                conn.commit()
                return {"data": notifications, "count": len(notifications)}

            if current_role in {"jefe_laboratorio", "laboratorio_tipificador"}:
                notifications = _fetch_laboratory_notifications(cur, current_role, limit=safe_limit)
                conn.commit()
                return {"data": notifications, "count": len(notifications)}

            return {"data": [], "count": 0}
    except Exception as e:
        logger.warning("Error fetching notification feed: %s", e)
        return {"data": [], "count": 0}
    finally:
        if 'conn' in locals() and conn:
            conn.close()


@app.get("/notifications/history")
async def get_notifications_history(request: Request, limit: int = 12):
    """Return resolved admin dashboard notifications as ticket history."""
    if not _has_database_url():
        return {"data": [], "count": 0}

    current_user_id = _extract_request_user_id(request)
    if not current_user_id:
        raise HTTPException(status_code=401, detail="Usuario no autenticado")

    safe_limit = max(1, min(int(limit or 12), 50))

    try:
        conn = _get_connection()
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            current_role = _normalize_role_name(_get_profile_role(cur, current_user_id))
            if current_role not in {"admin", "admin_general"}:
                return {"data": [], "count": 0}

            _sync_dashboard_notifications(cur, current_user_id)
            history = _fetch_dashboard_notification_history(cur, limit=safe_limit)
            conn.commit()
            return {"data": history, "count": len(history)}
    except Exception as e:
        logger.warning("Error fetching notification history: %s", e)
        return {"data": [], "count": 0}
    finally:
        if 'conn' in locals() and conn:
            conn.close()


@app.patch("/notifications/{notification_key}/acknowledge")
async def acknowledge_notification(notification_key: str, request: Request):
    """Mark a dashboard notification as acknowledged for the current role."""
    if not _has_database_url():
        raise HTTPException(status_code=400, detail="Database not configured")

    current_user_id = _extract_request_user_id(request)
    if not current_user_id:
        raise HTTPException(status_code=401, detail="Usuario no autenticado")

    try:
        conn = _get_connection()
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            current_role = _normalize_role_name(_get_profile_role(cur, current_user_id))
            notification = _get_dashboard_notification_by_key(cur, notification_key)
            if not notification:
                raise HTTPException(status_code=404, detail="Notificación no encontrada")
            if not _can_acknowledge_notification(current_role, notification):
                raise HTTPException(status_code=403, detail="No tienes permisos para gestionar esta notificación")
            row = _acknowledge_dashboard_notification(cur, notification_key, current_user_id)
            if not row:
                raise HTTPException(status_code=404, detail="Notificación no encontrada")
            conn.commit()
            return row
    except HTTPException:
        if 'conn' in locals() and conn:
            conn.rollback()
        raise
    except Exception as e:
        if 'conn' in locals() and conn:
            conn.rollback()
        logger.warning("Error acknowledging notification %s: %s", notification_key, e)
        raise HTTPException(status_code=500, detail="Error actualizando notificación")
    finally:
        if 'conn' in locals() and conn:
            conn.close()


@app.put("/roles/{role_id}")
async def update_role(role_id: str, payload: RoleUpdate):
    """Update a role using direct SQL for maximum reliability"""
    if not _has_database_url():
        raise HTTPException(status_code=400, detail="Database not configured")
    
    try:
        canonical_role_id = _normalize_role_name(role_id)
        conn = _get_connection()
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            cur.execute(
                """
                SELECT permissions
                FROM role_definitions
                WHERE role_id = %s
                LIMIT 1
                """,
                (canonical_role_id,),
            )
            existing_role_row = cur.fetchone()
            if not existing_role_row:
                raw_role_id = (role_id or "").strip().lower()
                if raw_role_id != canonical_role_id:
                    cur.execute(
                        """
                        SELECT permissions
                        FROM role_definitions
                        WHERE role_id = %s
                        LIMIT 1
                        """,
                        (raw_role_id,),
                    )
                    existing_role_row = cur.fetchone()

            # Prepare update map
            update_fields = []
            params = []
            
            if payload.label is not None:
                update_fields.append("label = %s")
                params.append(payload.label)
                
            if payload.description is not None:
                update_fields.append("description = %s")
                params.append(payload.description)
                
            if payload.permissions is not None:
                update_fields.append("permissions = %s")
                current_permissions = _normalize_permission_map((existing_role_row or {}).get("permissions") if isinstance(existing_role_row, dict) else None)
                incoming_permissions = _normalize_permission_map(payload.permissions.model_dump(exclude_unset=True))
                merged_permissions = _merge_permission_maps(current_permissions, incoming_permissions)
                merged_permissions = _grant_delete_to_oficina_tecnica(merged_permissions, canonical_role_id)
                # Ensure we serialize to JSON string for Postgres JSONB
                params.append(json.dumps(merged_permissions))
            
            if not update_fields:
                raise HTTPException(status_code=400, detail="No fields to update")
            
            update_fields.append("updated_at = NOW()")
            params.append(canonical_role_id)
            
            query = f"""
                UPDATE role_definitions 
                SET {', '.join(update_fields)} 
                WHERE role_id = %s
                RETURNING *
            """
            
            cur.execute(query, params)
            result = cur.fetchone()
            
            if not result:
                raw_role_id = (role_id or "").strip().lower()
                if raw_role_id != canonical_role_id:
                    params[-1] = raw_role_id
                    cur.execute(query, params)
                    result = cur.fetchone()
                if not result:
                    conn.rollback()
                    raise HTTPException(status_code=404, detail="Role not found")
                
            conn.commit()
            return dict(result)
            
    except Exception as e:
        print(f"Error updating role via SQL: {e}")
        if 'conn' in locals() and conn:
            conn.rollback()
        raise HTTPException(status_code=500, detail=str(e))
    finally:
        if 'conn' in locals() and conn:
            conn.close()


# --- User granular permission overrides ---

@app.get("/users/{user_id}/permissions-override")
async def get_user_permissions_override(user_id: str, request: Request):
    if not _has_database_url():
        raise HTTPException(status_code=400, detail="Database not configured")

    current_user_id = _extract_request_user_id(request)
    if not current_user_id:
        raise HTTPException(status_code=401, detail="Usuario no autenticado")

    try:
        conn = _get_connection()
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            current_role = _normalize_role_name(_get_profile_role(cur, current_user_id))
            is_admin = current_role in {"admin", "admin_general"}
            if not is_admin and current_user_id != user_id:
                raise HTTPException(status_code=403, detail="No autorizado para ver permisos de otro usuario")

            cur.execute(
                """
                SELECT user_id::text as user_id, enabled, permissions, updated_by::text as updated_by, updated_at
                FROM user_permission_overrides
                WHERE user_id = %s
                LIMIT 1
                """,
                (user_id,),
            )
            row = cur.fetchone()
            target_role = _normalize_role_name(_get_profile_role(cur, user_id))
            cur.execute(
                """
                SELECT permissions
                FROM role_definitions
                WHERE role_id = %s
                LIMIT 1
                """,
                (target_role,),
            )
            role_row = cur.fetchone()
            if not role_row:
                raw_role = (_get_profile_role(cur, user_id) or "").strip().lower()
                if raw_role and raw_role != target_role:
                    cur.execute(
                        """
                        SELECT permissions
                        FROM role_definitions
                        WHERE role_id = %s
                        LIMIT 1
                        """,
                        (raw_role,),
                    )
                    role_row = cur.fetchone()
            role_permissions = _normalize_permission_map((role_row or {}).get("permissions") if isinstance(role_row, dict) else None)
            role_permissions = _sanitize_permissions_for_role(target_role, role_permissions)
            role_permissions = _grant_delete_to_oficina_tecnica(role_permissions, target_role)
            if not row:
                effective_permissions = _merge_permission_maps(role_permissions, {})
                return {
                    "user_id": user_id,
                    "enabled": False,
                    "permissions": {},
                    "role_permissions": role_permissions,
                    "effective_permissions": effective_permissions,
                    "available_modules": _available_modules_for_role(target_role),
                    "updated_by": None,
                    "updated_at": None,
                }
            override_permissions = _compact_permission_override(role_permissions, _normalize_permission_map(row.get("permissions")))
            override_permissions = _sanitize_permissions_for_role(target_role, override_permissions)
            effective_permissions = _merge_permission_maps(role_permissions, override_permissions if bool(row.get("enabled")) else {})
            return {
                "user_id": row.get("user_id"),
                "enabled": bool(row.get("enabled")),
                "permissions": override_permissions,
                "role_permissions": role_permissions,
                "effective_permissions": effective_permissions,
                "available_modules": _available_modules_for_role(target_role),
                "updated_by": row.get("updated_by"),
                "updated_at": row.get("updated_at"),
            }
    except Exception as e:
        logger.warning(f"Error fetching permissions-override (DB might be down): {e}")
        # Retornar override vacío como fallback para no romper el dashboard
        return {
            "user_id": user_id,
            "enabled": False,
            "permissions": {},
            "role_permissions": {},
            "effective_permissions": {},
            "available_modules": _available_modules_for_role(None),
            "updated_by": None,
            "updated_at": None,
        }
    finally:
        if 'conn' in locals() and conn:
            conn.close()


@app.put("/users/{user_id}/permissions-override")
async def upsert_user_permissions_override(user_id: str, payload: UserPermissionOverrideUpdate, request: Request):
    if not _has_database_url():
        raise HTTPException(status_code=400, detail="Database not configured")

    current_user_id = _extract_request_user_id(request)
    if not current_user_id:
        raise HTTPException(status_code=401, detail="Usuario no autenticado")

    try:
        conn = _get_connection()
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            current_role = _normalize_role_name(_get_profile_role(cur, current_user_id))
            if current_role not in {"admin", "admin_general"}:
                raise HTTPException(status_code=403, detail="Solo administradores pueden editar permisos granulares")

            normalized_permissions = _normalize_permission_map(payload.permissions)
            target_role = _normalize_role_name(_get_profile_role(cur, user_id))
            cur.execute(
                """
                SELECT permissions
                FROM role_definitions
                WHERE role_id = %s
                LIMIT 1
                """,
                (target_role,),
            )
            role_row = cur.fetchone()
            role_permissions = _sanitize_permissions_for_role(target_role, _normalize_permission_map((role_row or {}).get("permissions") if isinstance(role_row, dict) else None))
            role_permissions = _grant_delete_to_oficina_tecnica(role_permissions, target_role)
            if _is_restricted_technical_role(target_role):
                forbidden_modules = [
                    module for module in (*_CONTROL_PERMISSION_MODULE_KEYS, *_RESTRICTED_TECHNICAL_MODULE_KEYS)
                    if module != "configuracion"
                    and any((normalized_permissions.get(module) or {}).values())
                ]
                if target_role == "tecnico_suelos" and any((normalized_permissions.get("configuracion") or {}).values()):
                    forbidden_modules = [
                        module for module in forbidden_modules if module != "configuracion"
                    ]
                if forbidden_modules:
                    raise HTTPException(
                        status_code=400,
                        detail="El rol técnico no puede recibir permisos de módulos de control.",
                    )
                if target_role != "tecnico_suelos":
                    normalized_permissions = _strip_control_permissions(normalized_permissions, target_role)
            normalized_permissions = _compact_permission_override(role_permissions, normalized_permissions) if payload.enabled else {}
            cur.execute(
                """
                INSERT INTO user_permission_overrides (user_id, enabled, permissions, updated_by, updated_at)
                VALUES (%s::uuid, %s, %s::jsonb, %s::uuid, NOW())
                ON CONFLICT (user_id)
                DO UPDATE SET
                    enabled = EXCLUDED.enabled,
                    permissions = EXCLUDED.permissions,
                    updated_by = EXCLUDED.updated_by,
                    updated_at = NOW()
                RETURNING user_id::text as user_id, enabled, permissions, updated_by::text as updated_by, updated_at
                """,
                (user_id, payload.enabled, json.dumps(normalized_permissions), current_user_id),
            )
            result = cur.fetchone()
            effective_permissions = _merge_permission_maps(role_permissions, _normalize_permission_map(result.get("permissions")) if bool(result.get("enabled")) else {})
            conn.commit()
            return {
                "user_id": result.get("user_id"),
                "enabled": bool(result.get("enabled")),
                "permissions": _normalize_permission_map(result.get("permissions")),
                "effective_permissions": effective_permissions,
                "available_modules": _available_modules_for_role(target_role),
                "updated_by": result.get("updated_by"),
                "updated_at": result.get("updated_at"),
            }
    except HTTPException:
        if 'conn' in locals() and conn:
            conn.rollback()
        raise
    except Exception as e:
        if 'conn' in locals() and conn:
            conn.rollback()
        raise HTTPException(status_code=500, detail=str(e))
    finally:
        if 'conn' in locals() and conn:
            conn.close()


@app.delete("/users/{user_id}/permissions-override")
async def clear_user_permissions_override(user_id: str, request: Request):
    if not _has_database_url():
        raise HTTPException(status_code=400, detail="Database not configured")

    current_user_id = _extract_request_user_id(request)
    if not current_user_id:
        raise HTTPException(status_code=401, detail="Usuario no autenticado")

    try:
        conn = _get_connection()
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            current_role = _normalize_role_name(_get_profile_role(cur, current_user_id))
            if current_role not in {"admin", "admin_general"}:
                raise HTTPException(status_code=403, detail="Solo administradores pueden limpiar permisos granulares")

            cur.execute("DELETE FROM user_permission_overrides WHERE user_id = %s::uuid", (user_id,))
            conn.commit()
            return {"success": True, "user_id": user_id}
    except HTTPException:
        if 'conn' in locals() and conn:
            conn.rollback()
        raise
    except Exception as e:
        if 'conn' in locals() and conn:
            conn.rollback()
        raise HTTPException(status_code=500, detail=str(e))
    finally:
        if 'conn' in locals() and conn:
            conn.close()


# --- Session Control Endpoints (Using Supabase REST API) ---

@app.post("/users/{user_id}/logout")
async def force_logout_user(user_id: str):
    """Force logout a user using Supabase REST API"""
    try:
        url = f"{_get_supabase_url()}/perfiles?id=eq.{user_id}"
        update_data = {"last_force_logout_at": datetime.utcnow().isoformat()}

        response = http_patch(
            url,
            headers=_get_supabase_headers(),
            json=update_data,
            timeout=5,
            request_name="supabase.perfiles.force_logout",
        )
        
        if response.status_code not in [200, 204]:
            raise HTTPException(status_code=500, detail=f"Error: {response.text}")
        
        return {"success": True, "message": "User session terminated"}
    except requests.RequestException as e:
        raise HTTPException(status_code=500, detail=str(e))


def _sync_heartbeat(user_id: str) -> dict:
    """Synchronous heartbeat logic — runs in a thread pool to avoid blocking the event loop."""
    headers = _get_supabase_headers()
    base_url = _get_supabase_url()

    # Check if user is active
    response = http_get(
        f"{base_url}/perfiles?id=eq.{user_id}&select=activo",
        headers=headers, timeout=5
    )
    if response.status_code != 200:
        return {"success": False, "error": "User not found"}

    data = response.json()
    if not data:
        return {"success": False, "error": "User not found"}

    if data[0].get("activo", True) is False:
        return {"success": False, "status": "inactive"}

    # Update last_seen_at
    update_response = http_patch(
        f"{base_url}/perfiles?id=eq.{user_id}",
        headers=headers,
        json={"last_seen_at": datetime.utcnow().isoformat()},
        timeout=5
    )
    if update_response.status_code not in [200, 204]:
        return {"success": False, "error": "Failed to update heartbeat"}

    return {"success": True, "status": "active"}


@app.post("/users/heartbeat")
async def user_heartbeat(payload: HeartbeatRequest):
    """Update user heartbeat using Supabase REST API (non-blocking)."""
    try:
        return await asyncio.to_thread(_sync_heartbeat, payload.user_id)
    except requests.RequestException as e:
        logger.exception("Heartbeat error for user_id=%s", payload.user_id)
        return {"success": False, "error": str(e)}


if __name__ == "__main__":
    import uvicorn

    uvicorn.run(app, host="0.0.0.0", port=8000)


# --- Centralized SQLAlchemy Event Listeners for Lab Essay Audit Logs ---
from sqlalchemy import event, text
from sqlalchemy.orm import Mapper

IGNORED_AUDIT_TABLES = {
    "cont_mat_organica_ensayos",
    "terrones_fino_grueso_ensayos",
    "azul_metileno_ensayos",
    "part_livianas_ensayos",
    "imp_organicas_ensayos",
    "sul_magnesio_ensayos",
    "angularidad_ensayos",
}

@event.listens_for(Mapper, "after_insert")
def audit_after_insert(mapper, connection, target):
    tablename = getattr(target, "__tablename__", "")
    if tablename and tablename != "auditoria" and tablename not in IGNORED_AUDIT_TABLES:
        if hasattr(target, "numero_ensayo") and hasattr(target, "numero_ot"):
            from app.auth import current_actor
            actor = current_actor.get() or {}
            user_id = actor.get("user_id")
            user_name = actor.get("user_name") or "Sistema"
            
            display_name = tablename.replace("_ensayos", "").replace("_", " ").title()
            
            try:
                connection.execute(
                    text(
                        """
                        INSERT INTO auditoria (user_id, user_name, action, module, details, severity, created_at)
                        VALUES (:user_id, :user_name, :action, :module, CAST(:details AS jsonb), 'info', NOW())
                        """
                    ),
                    {
                        "user_id": user_id,
                        "user_name": user_name,
                        "action": f"Creó ensayo de {display_name} {target.numero_ensayo}",
                        "module": "LABORATORIO",
                        "details": json.dumps({
                            "numero_ot": getattr(target, "numero_ot", None),
                            "muestra": getattr(target, "muestra", None),
                            "numero_ensayo": getattr(target, "numero_ensayo", None),
                            "id": getattr(target, "id", None)
                        }, ensure_ascii=False)
                    }
                )
            except Exception as e:
                logger.warning("SQLAlchemy audit insert log failed: %s", e)


@event.listens_for(Mapper, "after_update")
def audit_after_update(mapper, connection, target):
    tablename = getattr(target, "__tablename__", "")
    if tablename and tablename != "auditoria" and tablename not in IGNORED_AUDIT_TABLES:
        if hasattr(target, "numero_ensayo") and hasattr(target, "numero_ot"):
            from app.auth import current_actor
            actor = current_actor.get() or {}
            user_id = actor.get("user_id")
            user_name = actor.get("user_name") or "Sistema"
            
            display_name = tablename.replace("_ensayos", "").replace("_", " ").title()
            
            is_deleted = False
            deleted_at = getattr(target, "deleted_at", None)
            if deleted_at:
                is_deleted = True
                
            action_desc = "Eliminó" if is_deleted else "Actualizó"
            
            try:
                connection.execute(
                    text(
                        """
                        INSERT INTO auditoria (user_id, user_name, action, module, details, severity, created_at)
                        VALUES (:user_id, :user_name, :action, :module, CAST(:details AS jsonb), 'info', NOW())
                        """
                    ),
                    {
                        "user_id": user_id,
                        "user_name": user_name,
                        "action": f"{action_desc} ensayo de {display_name} {target.numero_ensayo}",
                        "module": "LABORATORIO",
                        "details": json.dumps({
                            "numero_ot": getattr(target, "numero_ot", None),
                            "muestra": getattr(target, "muestra", None),
                            "numero_ensayo": getattr(target, "numero_ensayo", None),
                            "id": getattr(target, "id", None)
                        }, ensure_ascii=False)
                    }
                )
            except Exception as e:
                logger.warning("SQLAlchemy audit update log failed: %s", e)


